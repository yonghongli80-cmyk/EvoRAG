"""Conservative structural edits that always produce a new `.xlsx` file."""

from __future__ import annotations

import re
from collections.abc import Mapping
from dataclasses import dataclass, replace
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from rag.spreadsheets.edit_models import (
    SpreadsheetStructureAction,
    SpreadsheetStructureEditRequest,
)

_INVALID_SHEET_NAME = re.compile(r"[\\/*?:\[\]]")
_MAX_ROW = 1_048_576
_MAX_COLUMN = 16_384

_STRUCTURE_ACTIONS = frozenset(
    {"insert_rows", "delete_rows", "insert_columns", "delete_columns"}
)


class WorkbookStructureError(Exception):
    def __init__(self, reason_code: str) -> None:
        super().__init__(reason_code)
        self.reason_code = reason_code


@dataclass(frozen=True)
class StructureChange:
    action: SpreadsheetStructureAction
    sheet_key: str | None
    sheet_name: str | None
    index: int | None
    count: int
    new_sheet_name: str | None
    affected_nonempty_cells: int


# ---------------------------------------------------------------------------
# Full-workbook invariants
#
# D-08: a structural edit used to be verified by comparing sheet names only, so
# a save that silently dropped images, validations, conditional formatting,
# comments, filters, or hyperlinks still produced a "successful" new version.
# The snapshot below records every structural surface we can compare across a
# save/reopen, and ``_validate_round_trip`` requires the reopened workbook to
# match the pre-edit snapshot (shifted exactly by the requested operation).
# ---------------------------------------------------------------------------


@dataclass(frozen=True)
class _SheetSnapshot:
    title: str
    state: str
    freeze_panes: str | None
    max_row: int
    max_column: int
    values: dict[tuple[int, int], object]
    styles: dict[tuple[int, int], tuple]
    merged: tuple[str, ...]
    tables: int
    images: int
    charts: int
    pivots: int
    data_validations: int
    conditional_formats: int
    comments: int
    hyperlinks: int
    autofilter: str | None
    protection: bool
    row_dimensions: dict[int, tuple[float | None, bool]]
    column_dimensions: dict[str, tuple[float | None, bool]]


@dataclass(frozen=True)
class _WorkbookSnapshot:
    sheet_names: tuple[str, ...]
    defined_names: tuple[str, ...]
    external_links: int
    lock_structure: bool
    sheets: tuple[_SheetSnapshot, ...]


def _color_key(color: Any) -> tuple | None:
    if color is None:
        return None
    return (
        getattr(color, "type", None),
        getattr(color, "rgb", None),
        getattr(color, "theme", None),
        getattr(color, "indexed", None),
    )


def _cell_style(cell) -> tuple:
    """A style signature that survives a save/reload (style ids may not)."""
    font = cell.font
    fill = cell.fill
    border = cell.border
    alignment = cell.alignment
    return (
        cell.number_format or "",
        bool(font.bold) if font else False,
        bool(font.italic) if font else False,
        _color_key(font.color) if font else None,
        fill.fill_type if fill else None,
        _color_key(fill.fgColor) if fill else None,
        (
            border.left.style,
            border.right.style,
            border.top.style,
            border.bottom.style,
        )
        if border
        else None,
        (alignment.horizontal, alignment.vertical) if alignment else None,
    )


def _snapshot_workbook(book) -> _WorkbookSnapshot:
    sheets = []
    for sheet in book.worksheets:
        values: dict[tuple[int, int], object] = {}
        styles: dict[tuple[int, int], tuple] = {}
        comments = 0
        hyperlinks = 0
        for row in sheet.iter_rows():
            for cell in row:
                coordinate = (cell.row, cell.column)
                if cell.value is not None:
                    values[coordinate] = cell.value
                if cell.has_style:
                    styles[coordinate] = _cell_style(cell)
                if cell.comment is not None:
                    comments += 1
                if cell.hyperlink is not None:
                    hyperlinks += 1
        comments += len(getattr(sheet, "_comments", ()) or ())
        sheets.append(
            _SheetSnapshot(
                title=sheet.title,
                state=sheet.sheet_state,
                freeze_panes=str(sheet.freeze_panes) if sheet.freeze_panes else None,
                max_row=sheet.max_row,
                max_column=sheet.max_column,
                values=values,
                styles=styles,
                merged=tuple(sorted(str(value) for value in sheet.merged_cells.ranges)),
                tables=len(getattr(sheet, "tables", {}) or {}),
                images=len(getattr(sheet, "_images", ()) or ()),
                charts=len(getattr(sheet, "_charts", ()) or ()),
                pivots=len(getattr(sheet, "_pivots", ()) or ()),
                data_validations=getattr(sheet.data_validations, "count", 0) or 0,
                conditional_formats=len(sheet.conditional_formatting),
                comments=comments,
                hyperlinks=hyperlinks,
                autofilter=getattr(getattr(sheet, "auto_filter", None), "ref", None),
                protection=bool(sheet.protection.sheet),
                row_dimensions={
                    index: (dimension.height, bool(dimension.hidden))
                    for index, dimension in sheet.row_dimensions.items()
                },
                column_dimensions={
                    key: (dimension.width, bool(dimension.hidden))
                    for key, dimension in sheet.column_dimensions.items()
                },
            )
        )
    return _WorkbookSnapshot(
        sheet_names=tuple(book.sheetnames),
        defined_names=tuple(sorted(str(name) for name in book.defined_names)),
        external_links=len(getattr(book, "_external_links", ()) or ()),
        lock_structure=bool(
            getattr(getattr(book, "security", None), "lockStructure", False)
        ),
        sheets=tuple(sheets),
    )


def _shifted_cells(
    cells: Mapping[tuple[int, int], Any],
    action: str,
    index: int,
    count: int,
) -> dict[tuple[int, int], Any]:
    """Return the cell map openpyxl's row/column shift is expected to produce."""
    result: dict[tuple[int, int], Any] = {}
    for (row, column), value in cells.items():
        if action == "insert_rows" and row >= index:
            row += count
        elif action == "delete_rows":
            if index <= row < index + count:
                continue
            if row >= index + count:
                row -= count
        elif action == "insert_columns" and column >= index:
            column += count
        elif action == "delete_columns":
            if index <= column < index + count:
                continue
            if column >= index + count:
                column -= count
        result[(row, column)] = value
    return result


def _validate_shifted_sheet(
    before: _SheetSnapshot,
    after: _SheetSnapshot,
    change: StructureChange,
) -> None:
    assert change.index is not None
    index, count = change.index, change.count
    expected_values = _shifted_cells(before.values, change.action, index, count)
    if after.values != expected_values:
        raise WorkbookStructureError("workbook_validation_failed")
    expected_styles = _shifted_cells(before.styles, change.action, index, count)
    if after.styles != expected_styles:
        raise WorkbookStructureError("workbook_validation_failed")
    # Every non-cell invariant (merges, tables, dimensions, protection, filters,
    # ...) must be preserved by the round trip.
    if replace(before, values={}, styles={}, max_row=0, max_column=0) != replace(
        after, values={}, styles={}, max_row=0, max_column=0
    ):
        raise WorkbookStructureError("workbook_validation_failed")
    occupied_rows = [row for row, _ in expected_values]
    occupied_columns = [column for _, column in expected_values]
    if after.max_row < max(occupied_rows, default=1):
        raise WorkbookStructureError("workbook_validation_failed")
    if after.max_column < max(occupied_columns, default=1):
        raise WorkbookStructureError("workbook_validation_failed")
    if change.action.startswith("insert_"):
        limit_row = before.max_row + (count if change.action == "insert_rows" else 0)
        limit_column = before.max_column + (
            count if change.action == "insert_columns" else 0
        )
    else:
        limit_row, limit_column = before.max_row, before.max_column
    if after.max_row > limit_row or after.max_column > limit_column:
        raise WorkbookStructureError("workbook_validation_failed")


def _validate_new_sheet(sheet: _SheetSnapshot) -> None:
    if (
        sheet.values
        or sheet.styles
        or sheet.merged
        or sheet.tables
        or sheet.images
        or sheet.charts
        or sheet.pivots
        or sheet.data_validations
        or sheet.conditional_formats
        or sheet.comments
        or sheet.hyperlinks
        or sheet.autofilter
        or sheet.freeze_panes
        or sheet.protection
        or sheet.row_dimensions
        or sheet.column_dimensions
    ):
        raise WorkbookStructureError("workbook_validation_failed")


def _validate_round_trip(
    before: _WorkbookSnapshot,
    after: _WorkbookSnapshot,
    change: StructureChange,
) -> None:
    """Reject a save that changed anything beyond the requested shift."""
    expected_names = (
        before.sheet_names + (change.new_sheet_name,)
        if change.action == "add_sheet"
        else before.sheet_names
    )
    if after.sheet_names != expected_names:
        raise WorkbookStructureError("workbook_validation_failed")
    if (
        before.defined_names != after.defined_names
        or before.external_links != after.external_links
        or before.lock_structure != after.lock_structure
    ):
        raise WorkbookStructureError("workbook_validation_failed")
    after_sheets = {sheet.title: sheet for sheet in after.sheets}
    for snapshot in before.sheets:
        current = after_sheets.get(snapshot.title)
        if current is None:
            raise WorkbookStructureError("workbook_validation_failed")
        if change.action in _STRUCTURE_ACTIONS and snapshot.title == change.sheet_name:
            _validate_shifted_sheet(snapshot, current, change)
        elif current != snapshot:
            # ``add_sheet`` must not disturb any pre-existing sheet; the same
            # identity check protects sheets the row/column edit did not target.
            raise WorkbookStructureError("workbook_validation_failed")
    if change.action == "add_sheet":
        if change.new_sheet_name is None:
            raise WorkbookStructureError("workbook_validation_failed")
        _validate_new_sheet(after_sheets[change.new_sheet_name])


def _load(path: Path):
    try:
        return load_workbook(
            path,
            data_only=False,
            read_only=False,
            keep_links=True,
        )
    except Exception as exc:
        raise WorkbookStructureError("workbook_open_failed") from exc


def _contains_formulas(book) -> bool:
    return any(
        cell.data_type == "f"
        for sheet in book.worksheets
        for row in sheet.iter_rows()
        for cell in row
    )


def _reject_unsupported_features(book) -> None:
    """Refuse features a row/column shift would move or detach (D-08)."""
    for sheet in book.worksheets:
        if getattr(sheet, "_images", ()):
            raise WorkbookStructureError("sheet_contains_images")
        if getattr(sheet.data_validations, "count", 0):
            raise WorkbookStructureError("sheet_contains_data_validations")
        if len(sheet.conditional_formatting):
            raise WorkbookStructureError("sheet_contains_conditional_formatting")
        if getattr(sheet, "auto_filter", None) is not None and sheet.auto_filter.ref:
            raise WorkbookStructureError("sheet_contains_autofilter")
        comments = len(getattr(sheet, "_comments", ()) or ())
        hyperlinks = 0
        for row in sheet.iter_rows():
            for cell in row:
                if cell.comment is not None or cell.hyperlink is not None:
                    if cell.comment is not None:
                        comments += 1
                    if cell.hyperlink is not None:
                        hyperlinks += 1
        if comments:
            raise WorkbookStructureError("sheet_contains_comments")
        if hyperlinks:
            raise WorkbookStructureError("sheet_contains_hyperlinks")


def _has_unsupported_dependencies(book) -> bool:
    if getattr(book, "_external_links", ()):
        return True
    if len(book.defined_names):
        return True
    return any(
        bool(sheet.tables)
        or bool(getattr(sheet, "_charts", ()))
        or bool(getattr(sheet, "_pivots", ()))
        for sheet in book.worksheets
    )


def _validate_row_or_column_edit(book, sheet) -> None:
    workbook_security = getattr(book, "security", None)
    sheet_protection = getattr(sheet, "protection", None)
    if bool(getattr(workbook_security, "lockStructure", False)) or bool(
        getattr(sheet_protection, "sheet", False)
    ):
        raise WorkbookStructureError("protected_workbook")
    if _contains_formulas(book):
        raise WorkbookStructureError("workbook_contains_formulas")
    if sheet.merged_cells.ranges:
        raise WorkbookStructureError("sheet_contains_merged_cells")
    _reject_unsupported_features(book)
    if _has_unsupported_dependencies(book):
        raise WorkbookStructureError("unsupported_workbook_dependencies")


def _validate_index(sheet, request: SpreadsheetStructureEditRequest) -> None:
    assert request.index is not None
    is_row = request.action.endswith("rows")
    maximum = _MAX_ROW if is_row else _MAX_COLUMN
    used = sheet.max_row if is_row else sheet.max_column
    final_index = request.index + request.count - 1
    if final_index > maximum:
        raise WorkbookStructureError("invalid_structure_index")
    if request.action.startswith("insert_"):
        if request.index > used + 1:
            raise WorkbookStructureError("invalid_structure_index")
    elif request.index > used or final_index > used:
        raise WorkbookStructureError("invalid_structure_index")


def _nonempty_delete_count(sheet, request: SpreadsheetStructureEditRequest) -> int:
    assert request.index is not None
    end = request.index + request.count - 1
    if request.action == "delete_rows":
        cells = (
            cell
            for row in sheet.iter_rows(
                min_row=request.index,
                max_row=end,
                min_col=1,
                max_col=sheet.max_column,
            )
            for cell in row
        )
    else:
        cells = (
            cell
            for row in sheet.iter_rows(
                min_row=1,
                max_row=sheet.max_row,
                min_col=request.index,
                max_col=end,
            )
            for cell in row
        )
    return sum(cell.value is not None for cell in cells)


def _validated_new_sheet_name(book, raw_name: str | None) -> str:
    name = (raw_name or "").strip()
    if (
        not name
        or len(name) > 31
        or _INVALID_SHEET_NAME.search(name)
        or name.casefold() in {existing.casefold() for existing in book.sheetnames}
    ):
        raise WorkbookStructureError("invalid_sheet_name")
    return name


def apply_structure_edit(
    source: Path,
    destination: Path,
    sheet_map: Mapping[str, str],
    request: SpreadsheetStructureEditRequest,
) -> StructureChange:
    """Apply exactly one safe structural operation and verify the saved workbook."""
    if source.suffix.lower() != ".xlsx" or destination.suffix.lower() != ".xlsx":
        raise WorkbookStructureError("unsupported_workbook_format")

    book = _load(source)
    try:
        before = _snapshot_workbook(book)
        if request.action == "add_sheet":
            new_sheet_name = _validated_new_sheet_name(book, request.new_sheet_name)
            book.create_sheet(new_sheet_name)
            change = StructureChange(
                action=request.action,
                sheet_key=None,
                sheet_name=None,
                index=None,
                count=1,
                new_sheet_name=new_sheet_name,
                affected_nonempty_cells=0,
            )
        else:
            assert request.sheet_key is not None
            sheet_name = sheet_map.get(request.sheet_key)
            if sheet_name is None or sheet_name not in book.sheetnames:
                raise WorkbookStructureError("unknown_sheet")
            sheet = book[sheet_name]
            _validate_row_or_column_edit(book, sheet)
            _validate_index(sheet, request)
            affected = 0
            if request.action.startswith("delete_"):
                affected = _nonempty_delete_count(sheet, request)
                if affected and not request.confirm_nonempty:
                    raise WorkbookStructureError(
                        "nonempty_delete_requires_confirmation"
                    )
            assert request.index is not None
            operations = {
                "insert_rows": sheet.insert_rows,
                "delete_rows": sheet.delete_rows,
                "insert_columns": sheet.insert_cols,
                "delete_columns": sheet.delete_cols,
            }
            operations[request.action](request.index, request.count)
            change = StructureChange(
                action=request.action,
                sheet_key=request.sheet_key,
                sheet_name=sheet_name,
                index=request.index,
                count=request.count,
                new_sheet_name=None,
                affected_nonempty_cells=affected,
            )

        destination.parent.mkdir(parents=True, exist_ok=True)
        book.save(destination)
    except WorkbookStructureError:
        destination.unlink(missing_ok=True)
        raise
    except Exception as exc:
        destination.unlink(missing_ok=True)
        raise WorkbookStructureError("workbook_validation_failed") from exc
    finally:
        book.close()

    reopened = None
    try:
        reopened = _load(destination)
        _validate_round_trip(before, _snapshot_workbook(reopened), change)
    except WorkbookStructureError:
        destination.unlink(missing_ok=True)
        raise
    except Exception as exc:
        destination.unlink(missing_ok=True)
        raise WorkbookStructureError("workbook_validation_failed") from exc
    finally:
        if reopened is not None:
            reopened.close()
    return change
