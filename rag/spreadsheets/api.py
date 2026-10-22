"""Versioned spreadsheet read and cell-edit routes."""
import asyncio
import sqlite3
from pathlib import Path

from fastapi import APIRouter, Depends, HTTPException, Query, Request, Response
from openpyxl.utils import get_column_letter, range_boundaries

from config import settings
from rag.api_security import require_doc_version_edit, require_doc_view, require_user
from rag.spreadsheets.audit import (
    get_spreadsheet_audit_detail,
    list_spreadsheet_audits,
)
from rag.spreadsheets.edit_models import (
    RequestMeta,
    SpreadsheetEditRequest,
    SpreadsheetEditUnchanged,
    SpreadsheetRestoreRequest,
    SpreadsheetStructureEditRequest,
)
from rag.spreadsheets.edit_service import (
    SpreadsheetEditService,
    SpreadsheetEditServiceError,
)
from rag.user_db import VersionConflict

router = APIRouter()

# A document version is readable through the row endpoint only while it is a
# live version.  Recycle/purge and failed/staging versions must not serve data
# even when the caller still holds a dataset path (D-02).
_READABLE_VERSION_STATUSES = frozenset({"current", "historical"})


@router.get("/admin/spreadsheet-changes")
def spreadsheet_change_audits(
    request: Request,
    user: dict = Depends(require_user),  # noqa: B008 - FastAPI dependency marker
    limit: int = Query(default=50, ge=1, le=100),
    cursor: str = "",
    document_id: int | None = None,
    actor_user_id: int | None = None,
    kb_id: str = "",
    operation: str = "",
    status: str = "",
    risk: str = "",
    created_after: float | None = None,
    created_before: float | None = None,
):
    try:
        return list_spreadsheet_audits(
            request.app.state.user_db,
            user,
            limit=limit,
            cursor=cursor,
            document_id=document_id,
            actor_user_id=actor_user_id,
            kb_id=kb_id,
            operation=operation,
            status=status,
            risk=risk,
            created_after=created_after,
            created_before=created_before,
        )
    except PermissionError as exc:
        raise HTTPException(status_code=403, detail="audit management permission required") from exc
    except ValueError as exc:
        raise HTTPException(status_code=422, detail="invalid audit cursor") from exc


@router.get("/admin/spreadsheet-changes/{session_id}")
def spreadsheet_change_audit_detail(
    session_id: int,
    request: Request,
    user: dict = Depends(require_user),  # noqa: B008 - FastAPI dependency marker
):
    detail = get_spreadsheet_audit_detail(
        request.app.state.user_db, user, session_id
    )
    if detail is None:
        raise HTTPException(status_code=404, detail="spreadsheet audit not found")
    return detail


def _edit_service(request: Request) -> SpreadsheetEditService:
    return SpreadsheetEditService(
        request.app.state.user_db,
        Path(
            getattr(
                request.app.state,
                "document_versions_root",
                settings.document_versions_root,
            )
        ),
        Path(
            getattr(
                request.app.state,
                "structured_datasets_root",
                settings.structured_datasets_root,
            )
        ),
    )


@router.post("/documents/{document_id}/spreadsheet-edits")
async def create_spreadsheet_edit(
    document_id: int,
    body: SpreadsheetEditRequest,
    request: Request,
    response: Response,
    user: dict = Depends(require_user),  # noqa: B008 - FastAPI dependency marker
):
    meta = RequestMeta(
        ip_address=request.client.host if request.client else "",
        user_agent=request.headers.get("user-agent", ""),
    )
    try:
        result = await asyncio.to_thread(
            _edit_service(request).save,
            document_id,
            user,
            body,
            meta,
        )
    except VersionConflict as exc:
        raise HTTPException(
            status_code=409,
            detail={"code": str(exc)},
        ) from exc
    except SpreadsheetEditServiceError as exc:
        raise HTTPException(
            status_code=exc.status_code,
            detail={"code": exc.reason_code},
        ) from exc
    if isinstance(result, SpreadsheetEditUnchanged):
        response.status_code = 200
    else:
        response.status_code = 202
    return result.model_dump()


@router.post(
    "/documents/{document_id}/spreadsheet-structure-edits",
    status_code=202,
)
async def create_spreadsheet_structure_edit(
    document_id: int,
    body: SpreadsheetStructureEditRequest,
    request: Request,
    user: dict = Depends(require_user),  # noqa: B008 - FastAPI dependency marker
):
    meta = RequestMeta(
        ip_address=request.client.host if request.client else "",
        user_agent=request.headers.get("user-agent", ""),
    )
    try:
        result = await asyncio.to_thread(
            _edit_service(request).save_structure,
            document_id,
            user,
            body,
            meta,
        )
    except VersionConflict as exc:
        raise HTTPException(
            status_code=409,
            detail={"code": str(exc)},
        ) from exc
    except SpreadsheetEditServiceError as exc:
        raise HTTPException(
            status_code=exc.status_code,
            detail={"code": exc.reason_code},
        ) from exc
    return result.model_dump()


@router.get("/documents/{document_id}/spreadsheet-edits/{session_id}")
async def get_spreadsheet_edit_status(
    document_id: int,
    session_id: int,
    request: Request,
    user: dict = Depends(require_user),  # noqa: B008 - FastAPI dependency marker
):
    db = request.app.state.user_db
    session = await asyncio.to_thread(db.get_spreadsheet_edit_session, session_id)
    if session is None or session["document_id"] != document_id:
        raise HTTPException(status_code=404, detail="spreadsheet edit not found")
    document = await asyncio.to_thread(db.get_document_permission_by_id, document_id)
    if document is None or not (
        user.get("is_admin")
        or session["actor_user_id"] == user["id"]
        or document["owner_id"] == user["id"]
    ):
        raise HTTPException(status_code=404, detail="spreadsheet edit not found")
    return {
        "session_id": session["id"],
        "version_id": session["new_version_id"],
        "status": session["status"],
        "error_code": session["error_code"],
    }


@router.post(
    "/documents/{document_id}/versions/{source_version_id}/restore",
    status_code=202,
)
async def restore_spreadsheet_version(
    document_id: int,
    source_version_id: int,
    body: SpreadsheetRestoreRequest,
    request: Request,
    user: dict = Depends(require_user),  # noqa: B008 - FastAPI dependency marker
):
    meta = RequestMeta(
        ip_address=request.client.host if request.client else "",
        user_agent=request.headers.get("user-agent", ""),
    )
    try:
        result = await asyncio.to_thread(
            _edit_service(request).restore,
            document_id,
            source_version_id,
            user,
            body,
            meta,
        )
    except VersionConflict as exc:
        raise HTTPException(
            status_code=409,
            detail={"code": str(exc)},
        ) from exc
    except SpreadsheetEditServiceError as exc:
        raise HTTPException(
            status_code=exc.status_code,
            detail={"code": exc.reason_code},
        ) from exc
    return result.model_dump()

@router.get("/documents/{document_id}/versions/{version_id}/sheets/{sheet_key}/rows")
def rows(
    document_id: int,
    version_id: int,
    sheet_key: str,
    request: Request,
    user: dict = Depends(require_user),  # noqa: B008 - FastAPI dependency marker
    after_row: int = 0,
    limit: int = Query(default=100, ge=1, le=500),
):
    db = request.app.state.user_db
    document = db.get_document_permission_by_id(document_id)
    if document is None:
        raise HTTPException(404, "document not found")
    require_doc_view(db, document["doc_name"], document["kb_id"], user)
    version = db.get_document_version_by_id(version_id)
    dataset = db.get_structured_dataset_for_version(version_id)
    if (
        version is None
        or dataset is None
        or version["document_id"] != document_id
        or dataset["document_id"] != document_id
        or dataset["status"] != "ready"
        or version["status"] not in _READABLE_VERSION_STATUSES
    ):
        raise HTTPException(404, "spreadsheet version not found")
    root = Path(
        getattr(
            request.app.state,
            "structured_datasets_root",
            settings.structured_datasets_root,
        )
    )
    path = (root / dataset["storage_key"]).resolve()
    if root.resolve() not in path.parents:
        raise HTTPException(404, "spreadsheet version not found")
    connection = sqlite3.connect(f"file:{path.as_posix()}?mode=ro", uri=True)
    try:
        sheet = connection.execute("SELECT table_name FROM sheet_manifest WHERE sheet_key = ?", (sheet_key,)).fetchone()
        columns = connection.execute(
            "SELECT column_key, ordinal, display_name, inferred_type "
            "FROM column_manifest WHERE sheet_key = ? ORDER BY ordinal",
            (sheet_key,),
        ).fetchall()
        if sheet is None:
            raise HTTPException(404, "sheet not found")
        keys = [row[0] for row in columns]
        quoted = ", ".join(f'"{key}"' for key in keys)
        result = connection.execute(f'SELECT __row_number, {quoted} FROM "{sheet[0]}" WHERE __row_number > ? ORDER BY __row_number LIMIT ?', (after_row, limit)).fetchall()
        non_editable = {
            f"{get_column_letter(int(str(row[1]).removeprefix('c_')))}{row[0]}"
            for row in connection.execute(
                "SELECT row_number, column_key FROM formula_cells WHERE sheet_key = ?",
                (sheet_key,),
            ).fetchall()
        }
        for merged_row in connection.execute(
            "SELECT range_ref FROM merged_ranges WHERE sheet_key = ?", (sheet_key,)
        ).fetchall():
            min_col, min_row, max_col, max_row = range_boundaries(merged_row[0])
            for row_number in range(min_row, max_row + 1):
                for column_number in range(min_col, max_col + 1):
                    if (column_number, row_number) != (min_col, min_row):
                        non_editable.add(
                            f"{get_column_letter(column_number)}{row_number}"
                        )
        try:
            require_doc_version_edit(db, document_id, user)
            can_edit = document.get("current_version_id") == version_id
        except HTTPException:
            can_edit = False
        return {
            "rows": [
                {"row_number": row[0], "values": list(row[1:])}
                for row in result
            ],
            "columns": [
                {
                    "column_key": row[0],
                    "ordinal": row[1],
                    "display_name": row[2],
                    "value_type": row[3],
                }
                for row in columns
            ],
            "can_edit": can_edit,
            "current_version_id": document.get("current_version_id"),
            "non_editable_cells": sorted(non_editable),
        }
    finally:
        connection.close()
