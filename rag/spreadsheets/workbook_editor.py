"""Apply cell-only spreadsheet patches while preserving workbook structure."""

from __future__ import annotations

import math
from collections.abc import Mapping
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string, coordinate_from_string
from openpyxl.utils.exceptions import CellCoordinatesException

from rag.spreadsheets.edit_models import CellPatch, CellValueType


class WorkbookEditError(Exception):
    def __init__(self, reason_code: str) -> None:
        super().__init__(reason_code)
        self.reason_code = reason_code


@dataclass(frozen=True)
class CellChange:
    sheet_key: str
    sheet_name: str
    cell_address: str
    old_value_type: CellValueType
    old_value: object | None
    new_value_type: CellValueType
    new_value: object | None


def _value_type(value: object | None) -> CellValueType:
    if value is None:
        return "blank"
    if type(value) is bool:
        return "boolean"
    if isinstance(value, (date, datetime)):
        return "date"
    if type(value) in {int, float}:
        return "number"
    return "text"


def _converted_value(patch: CellPatch) -> object | None:
    if patch.value_type == "blank":
        return None
    if patch.value_type != "date":
        return patch.value
    raw = str(patch.value)
    if "T" in raw or " " in raw:
        return datetime.fromisoformat(raw)
    return date.fromisoformat(raw)


def _validated_address(raw_address: str) -> str:
    try:
        column, row = coordinate_from_string(raw_address)
        column_index = column_index_from_string(column)
    except (CellCoordinatesException, TypeError, ValueError):
        raise WorkbookEditError("invalid_cell_address") from None
    canonical = f"{column.upper()}{row}"
    if (
        raw_address.upper() != canonical
        or row < 1
        or row > 1_048_576
        or column_index < 1
        or column_index > 16_384
    ):
        raise WorkbookEditError("invalid_cell_address")
    return canonical


def _validate_patch_value(patch: CellPatch) -> None:
    value = patch.value
    if isinstance(value, str) and value.startswith("="):
        raise WorkbookEditError("formula_injection")
    if patch.value_type == "blank":
        valid = value is None
    elif patch.value_type == "text":
        valid = type(value) is str and len(value) <= 32_767
    elif patch.value_type == "number":
        valid = type(value) in {int, float} and math.isfinite(value)
    elif patch.value_type == "boolean":
        valid = type(value) is bool
    elif patch.value_type == "date":
        valid = type(value) is str
        if valid:
            try:
                converted = _converted_value(patch)
                valid = not isinstance(converted, datetime) or converted.tzinfo is None
            except ValueError:
                valid = False
    else:
        valid = False
    if not valid:
        raise WorkbookEditError("invalid_cell_value")


def _saved_value_matches(actual: object | None, expected: object | None) -> bool:
    if isinstance(expected, date) and not isinstance(expected, datetime):
        return isinstance(actual, datetime) and actual.date() == expected
    return type(actual) is type(expected) and actual == expected


def _snapshot_workbook(book) -> dict[str, Any]:
    sheets: dict[str, Any] = {}
    for sheet in book.worksheets:
        formulas: dict[str, str] = {}
        styles: dict[str, int] = {}
        for row in sheet.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    formulas[cell.coordinate] = str(cell.value)
                if cell.has_style:
                    styles[cell.coordinate] = cell.style_id
        sheets[sheet.title] = {
            "state": sheet.sheet_state,
            "freeze_panes": str(sheet.freeze_panes) if sheet.freeze_panes else None,
            "merged_ranges": tuple(str(value) for value in sheet.merged_cells.ranges),
            "formulas": formulas,
            "styles": styles,
            "row_dimensions": {
                index: (dimension.height, dimension.hidden)
                for index, dimension in sheet.row_dimensions.items()
            },
            "column_dimensions": {
                index: (dimension.width, dimension.hidden)
                for index, dimension in sheet.column_dimensions.items()
            },
            "protection": bool(sheet.protection.sheet),
        }
    return {
        "sheet_names": tuple(book.sheetnames),
        "lock_structure": bool(book.security.lockStructure),
        "sheets": sheets,
    }


def apply_cell_patches(
    source: Path,
    destination: Path,
    sheet_map: Mapping[str, str],
    patches: tuple[CellPatch, ...],
) -> tuple[CellChange, ...]:
    """Write typed cell values to a new `.xlsx` and verify no structure drift."""
    if source.suffix.lower() != ".xlsx" or destination.suffix.lower() != ".xlsx":
        raise WorkbookEditError("unsupported_workbook_format")
    if not patches or len(patches) > 500:
        raise WorkbookEditError("invalid_patch_count")

    try:
        book = load_workbook(
            source,
            data_only=False,
            read_only=False,
            keep_links=True,
        )
    except Exception as exc:
        raise WorkbookEditError("workbook_open_failed") from exc

    before = _snapshot_workbook(book)
    changes: list[CellChange] = []
    expected: list[tuple[str, str, object | None]] = []
    try:
        seen: set[tuple[str, str]] = set()
        for patch in patches:
            sheet_name = sheet_map.get(patch.sheet_key)
            if sheet_name is None or sheet_name not in book.sheetnames:
                raise WorkbookEditError("unknown_sheet")
            address = _validated_address(patch.cell)
            unique_key = (patch.sheet_key.casefold(), address)
            if unique_key in seen:
                raise WorkbookEditError("duplicate_cell_patch")
            seen.add(unique_key)
            _validate_patch_value(patch)
            sheet = book[sheet_name]
            for merged_range in sheet.merged_cells.ranges:
                if address in merged_range and address != merged_range.start_cell.coordinate:
                    raise WorkbookEditError("merged_cell_not_editable")
            cell = sheet[address]
            if cell.data_type == "f":
                raise WorkbookEditError("formula_cell_not_editable")
            if sheet.protection.sheet and cell.protection.locked:
                raise WorkbookEditError("protected_cell")
            new_value = _converted_value(patch)
            changes.append(
                CellChange(
                    sheet_key=patch.sheet_key,
                    sheet_name=sheet_name,
                    cell_address=cell.coordinate,
                    old_value_type=_value_type(cell.value),
                    old_value=cell.value,
                    new_value_type=patch.value_type,
                    new_value=new_value,
                )
            )
            cell.value = new_value
            expected.append((sheet_name, cell.coordinate, new_value))

        destination.parent.mkdir(parents=True, exist_ok=True)
        book.save(destination)
        reopened = load_workbook(
            destination,
            data_only=False,
            read_only=False,
            keep_links=True,
        )
        if _snapshot_workbook(reopened) != before:
            raise WorkbookEditError("workbook_validation_failed")
        for sheet_name, address, expected_value in expected:
            if not _saved_value_matches(
                reopened[sheet_name][address].value, expected_value
            ):
                raise WorkbookEditError("workbook_validation_failed")
    except WorkbookEditError:
        destination.unlink(missing_ok=True)
        raise
    except Exception as exc:
        destination.unlink(missing_ok=True)
        raise WorkbookEditError("workbook_validation_failed") from exc
    return tuple(changes)
