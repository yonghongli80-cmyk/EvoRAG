import os
import shutil
import tempfile
import warnings
from pathlib import Path

from config import settings
from rag.cleaner import clean_text, detect_and_decode


def load(file_path: str) -> str:
    """Join normalized extraction units for legacy callers.

    The unified format-capability registry routes clean files through the
    extractors; legacy per-format loaders remain as a lenient fallback so
    existing callers keep their historical behavior. Strict signature/MIME
    enforcement lives in the material staging boundary.
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"文件不存在: {file_path}")
    if os.path.getsize(file_path) > settings.max_upload_bytes:
        raise ValueError(f"文件大小超过 {settings.max_upload_bytes} 字节限制")
    path = Path(file_path)
    ext = path.suffix.lower()
    from rag.materials.formats import build_default_format_registry
    from rag.materials.extraction import extract_material

    with path.open("rb") as source:
        signature = source.read(16)
    resolution = build_default_format_registry().resolve(path.name, None, signature)
    if resolution.can_index:
        material = extract_material(path, version_id=0, resolution=resolution)
        return "\n\n".join(unit.text for unit in material.units if unit.text)
    if ext == ".pdf":
        return load_pdf(file_path)
    if ext == ".docx":
        return load_docx(file_path)
    if ext == ".xlsx":
        return load_excel(file_path)
    with open(file_path, "rb") as f:
        raw = f.read(settings.max_upload_bytes + 1)
    if len(raw) > settings.max_upload_bytes:
        raise ValueError(f"文件大小超过 {settings.max_upload_bytes} 字节限制")
    text = detect_and_decode(raw)
    return clean_text(text)


def _supported_extensions() -> set[str]:
    """Compatibility set derived from the unified format registry."""
    from rag.materials.formats import build_default_format_registry

    registry = build_default_format_registry()
    extensions: set[str] = set()
    for capability in registry._by_key.values():
        if capability.extract_text or capability.structured_query or capability.semantic_index:
            extensions.update(f".{ext}" for ext in capability.extensions)
    return extensions


SUPPORTED_EXTENSIONS = _supported_extensions()


def _ensure_java_on_path():
    """Put Java on PATH if not already accessible (for OpenDataLoader PDF)."""
    import shutil

    if shutil.which("java"):
        return
    candidates = [
        os.environ.get("JAVA_HOME", ""),
        r"C:\Program Files\Amazon Corretto\jdk17.0.19_10",
        r"C:\Program Files\Java\jdk-17",
        r"C:\Program Files\Eclipse Adoptium\jdk-17",
    ]
    for base in candidates:
        if not base:
            continue
        java_bin = os.path.join(base, "bin")
        if os.path.isfile(os.path.join(java_bin, "java.exe")):
            os.environ["PATH"] = java_bin + os.pathsep + os.environ["PATH"]
            return


def load_pdf(file_path: str) -> str:
    """Extract text from PDF using OpenDataLoader (converts to Markdown)."""
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"文件不存在: {file_path}")

    import opendataloader_pdf

    _ensure_java_on_path()
    tmp_dir = tempfile.mkdtemp()
    try:
        opendataloader_pdf.convert(
            input_path=[file_path],
            output_dir=tmp_dir,
            format="markdown",
        )
        stem = Path(file_path).stem
        md_path = Path(tmp_dir) / f"{stem}.md"
        if md_path.exists():
            return md_path.read_text(encoding="utf-8")
        warnings.warn(f"OpenDataLoader PDF 未生成输出文件: {md_path}")
        return ""
    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)


def load_docx(file_path: str) -> str:
    """Extract text from DOCX using python-docx (paragraphs + tables in order)."""
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"文件不存在: {file_path}")
    from docx import Document
    from docx.oxml.ns import qn

    doc = Document(file_path)
    parts = []
    for elem in doc.element.body:
        if elem.tag == qn("w:p"):
            texts = [t.text for t in elem.iter(qn("w:t")) if t.text]
            line = "".join(texts).strip()
            if line:
                parts.append(line)
        elif elem.tag == qn("w:tbl"):
            for row in elem.iter(qn("w:tr")):
                cells = []
                for cell in row.iter(qn("w:tc")):
                    texts = [t.text for t in cell.iter(qn("w:t")) if t.text]
                    cells.append("".join(texts).strip())
                line = " | ".join(cells)
                if line.strip():
                    parts.append(line)
    return "\n".join(parts)


def load_excel(file_path: str) -> str:
    """Read .xlsx file, convert all sheets to text."""
    from openpyxl import load_workbook

    try:
        wb = load_workbook(file_path, read_only=True, data_only=True)
    except Exception as e:
        raise ValueError(f"无法读取 Excel 文件: {file_path} ({e})")
    parts = []
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        parts.append(f"=== Sheet: {sheet} ===")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            parts.append("\t".join(cells))
    wb.close()
    return "\n".join(parts)
