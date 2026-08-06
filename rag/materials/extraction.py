"""Normalized extraction dispatch and the extractor registry."""

from __future__ import annotations

from pathlib import Path

from rag.materials.models import (
    CitationLocation,
    ExtractedMaterial,
    ExtractedUnit,
    FormatResolution,
)
from rag.materials.pdf_ocr import LocalOcrEngine, extract_pdf
from rag.materials.presentation import extract_presentation


def _extract_text_file(path: Path, version_id: int) -> ExtractedMaterial:
    from rag.cleaner import clean_text, detect_and_decode

    raw = path.read_bytes()
    text = clean_text(detect_and_decode(raw))
    units = (
        ExtractedUnit(
            unit_key="text",
            text=text,
            location=CitationLocation(kind="paragraph", value="1"),
            source="native_text",
            confidence=1.0,
            metadata={},
        ),
    )
    return ExtractedMaterial(version_id, "text-v1", units, ())


def _extract_docx(path: Path, version_id: int) -> ExtractedMaterial:
    from docx import Document
    from docx.oxml.ns import qn

    document = Document(str(path))
    parts: list[str] = []
    for element in document.element.body:
        if element.tag == qn("w:p"):
            line = "".join(t.text for t in element.iter(qn("w:t")) if t.text).strip()
            if line:
                parts.append(line)
        elif element.tag == qn("w:tbl"):
            for row in element.iter(qn("w:tr")):
                cells = [
                    "".join(t.text for t in cell.iter(qn("w:t")) if t.text).strip()
                    for cell in row.iter(qn("w:tc"))
                ]
                line = " | ".join(cells)
                if line.strip():
                    parts.append(line)
    text = "\n".join(parts)
    units = (
        ExtractedUnit(
            unit_key="body",
            text=text,
            location=CitationLocation(kind="paragraph", value="1"),
            source="native_text",
            confidence=1.0,
            metadata={},
        ),
    )
    return ExtractedMaterial(version_id, "docx-v1", units, ())


def _extract_xlsx(path: Path, version_id: int) -> ExtractedMaterial:
    from openpyxl import load_workbook

    workbook = load_workbook(str(path), read_only=True, data_only=True)
    try:
        parts: list[str] = []
        for sheet in workbook.sheetnames:
            worksheet = workbook[sheet]
            parts.append(f"=== Sheet: {sheet} ===")
            for row in worksheet.iter_rows(values_only=True):
                cells = [str(c) if c is not None else "" for c in row]
                parts.append("\t".join(cells))
        text = "\n".join(parts)
    finally:
        workbook.close()
    units = (
        ExtractedUnit(
            unit_key="sheet",
            text=text,
            location=CitationLocation(kind="cell_range", value="sheet"),
            source="native_text",
            confidence=1.0,
            metadata={},
        ),
    )
    return ExtractedMaterial(version_id, "xlsx-v1", units, ())


def extract_material(
    path: Path,
    version_id: int,
    resolution: FormatResolution,
    ocr: LocalOcrEngine | None = None,
) -> ExtractedMaterial:
    """Dispatch to the format-specific extractor for one resolved material."""
    capability = resolution.capability
    if capability is None:
        raise ValueError(f"unsupported material: {resolution.reason_code}")
    key = capability.key
    if key == "pdf":
        return extract_pdf(path, version_id, ocr=ocr)
    if key == "pptx":
        return extract_presentation(path, version_id)
    if key in {"txt", "md", "csv"}:
        return _extract_text_file(path, version_id)
    if key == "docx":
        return _extract_docx(path, version_id)
    if key == "xlsx":
        return _extract_xlsx(path, version_id)
    raise ValueError(f"unsupported extraction for format: {key}")


class MaterialExtractorRegistry:
    """Map format capability keys to their extractors."""

    def __init__(self, extractors: dict[str, callable] | None = None):
        self._extractors = extractors or {}

    def register(self, key: str, extractor: callable) -> None:
        self._extractors[key] = extractor

    def extract(self, key: str, path: Path, version_id: int) -> ExtractedMaterial:
        extractor = self._extractors.get(key)
        if extractor is None:
            raise ValueError(f"no extractor for format: {key}")
        return extractor(path, version_id)
