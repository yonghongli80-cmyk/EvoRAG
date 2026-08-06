"""Page-aware ``.pptx`` extraction: slide text, speaker notes, and tables.

Embedded images are counted as metadata only; they are never OCR'd,
captioned, embedded, or sent to a vision model.
"""

from __future__ import annotations

from pathlib import Path

from rag.materials.models import CitationLocation, ExtractedMaterial, ExtractedUnit


def _shape_text(shape) -> str:
    if hasattr(shape, "has_text_frame") and shape.has_text_frame:
        lines = []
        for paragraph in shape.text_frame.paragraphs:
            line = "".join(run.text for run in paragraph.runs).strip()
            if line:
                lines.append(line)
        return "\n".join(lines)
    return ""


def _table_text(table) -> str:
    lines = []
    for row in table.rows:
        cells = [cell.text.strip() for cell in row.cells]
        lines.append(" | ".join(cells))
    return "\n".join(lines)


def extract_presentation(path: Path, version_id: int) -> ExtractedMaterial:
    from pptx import Presentation

    presentation = Presentation(str(path))
    units: list[ExtractedUnit] = []
    for slide_no, slide in enumerate(presentation.slides, start=1):
        parts: list[str] = []
        image_count = 0

        if slide.shapes.title is not None and slide.shapes.title.text.strip():
            parts.append(slide.shapes.title.text.strip())

        for shape in slide.shapes:
            shape_type = str(getattr(shape, "shape_type", ""))
            if "PICTURE" in shape_type.upper():
                image_count += 1
                continue
            text = _shape_text(shape)
            if text:
                parts.append(text)
            if getattr(shape, "has_table", False) and shape.has_table:
                table_text = _table_text(shape.table)
                if table_text:
                    parts.append(table_text)

        if slide.has_notes_slide:
            notes = (slide.notes_slide.notes_text_frame.text or "").strip()
            if notes:
                parts.append(notes)

        units.append(
            ExtractedUnit(
                unit_key=f"slide:{slide_no}",
                text="\n".join(parts),
                location=CitationLocation(kind="slide", value=str(slide_no)),
                source="native_text",
                confidence=1.0,
                metadata={"image_count": str(image_count)},
            )
        )
    return ExtractedMaterial(
        version_id=version_id,
        parser_schema_version="pptx-v1",
        units=tuple(units),
        warnings=(),
    )
