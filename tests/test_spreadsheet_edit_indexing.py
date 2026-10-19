"""End-to-end worker activation for immutable spreadsheet edits."""

from dataclasses import replace

from openpyxl import Workbook

from rag.document_version_store import final_version_path
from rag.materials.formats import build_default_format_registry
from rag.materials.worker import WorkerServices, process_index_job
from rag.models import Chunk
from rag.spreadsheets.edit_models import (
    CellPatch,
    RequestMeta,
    SpreadsheetEditRequest,
    SpreadsheetStructureEditRequest,
)
from rag.spreadsheets.edit_service import SpreadsheetEditService
from rag.user_db import UserDB


class RecordingEmbedder:
    def __init__(self):
        self.calls: list[list[str]] = []

    def __call__(self, texts):
        self.calls.append(list(texts))
        return [[float(len(text))] * 4 for text in texts]


class MemoryVectorStore:
    def __init__(self):
        self.versions = {}

    def add_units(self, kb_id, chunks, vectors, **kwargs):
        self.versions.setdefault(kwargs["document_version_id"], []).extend(chunks)

    def copy_units_to_version(
        self, kb_id, source_version_id, target_version_id, document_id, version_no, unit_keys,
        dataset_id=None,
    ):
        requested = set(unit_keys)
        source = {
            chunk.metadata.get("unit_key"): chunk
            for chunk in self.versions.get(source_version_id, [])
        }
        if not requested.issubset(source):
            return 0
        self.versions.setdefault(target_version_id, []).extend(
            replace(
                source[key],
                document_version_id=target_version_id,
                version_no=version_no,
                is_current=True,
                dataset_id=dataset_id or source[key].dataset_id,
            )
            for key in unit_keys
        )
        return len(unit_keys)

    def delete_version(self, kb_id, version_id):
        self.versions.pop(version_id, None)

    def mark_current(self, kb_id, document_id, version_id):
        return None


def test_edit_job_builds_complete_dataset_and_reuses_unchanged_vectors(tmp_path):
    db = UserDB(str(tmp_path / "users.sqlite3"))
    db.create_material_schema()
    owner_id = db.create_user("owner", "Password-123!")
    owner = db.get_user_by_id(owner_id)
    document_id = db.create_document_permission("prices.xlsx", "rag_docs", owner_id)
    versions_root = tmp_path / "versions"
    datasets_root = tmp_path / "datasets"
    source = final_version_path(versions_root, document_id, 1, "prices.xlsx")
    source.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.active.title = "Prices"
    workbook.active.append(["product", "price"])
    workbook.active.append(["Phone", 6800])
    workbook.save(source)
    v1 = db.create_document_version(
        document_id, "prices.xlsx", str(source), "hash-1", owner_id
    )
    vectors = MemoryVectorStore()
    embedder = RecordingEmbedder()
    services = WorkerServices(
        db=db,
        versions_root=versions_root,
        structured_datasets_root=datasets_root,
        registry=build_default_format_registry(),
        vector_store=vectors,
        embedder=embedder,
    )

    first = process_index_job(
        {"id": 1, "document_id": document_id, "version_id": v1}, services
    )
    assert first.status == "done"
    assert db.get_structured_dataset_for_version(v1)["status"] == "ready"

    accepted = SpreadsheetEditService(
        db, versions_root, datasets_root
    ).save(
        document_id,
        owner,
        SpreadsheetEditRequest(
            base_version_id=v1,
            changes=(
                CellPatch(
                    sheet_key="sheet_0001",
                    cell="B2",
                    value_type="number",
                    value=6500,
                ),
            ),
            reason="price correction",
        ),
        RequestMeta(),
    )
    embedder.calls.clear()

    second = process_index_job(
        {
            "id": 2,
            "document_id": document_id,
            "version_id": accepted.version_id,
        },
        services,
    )

    assert second.status == "done"
    assert second.reused_units == 2
    assert second.embedded_units == 1
    assert len(embedder.calls) == 1
    assert "6500" in embedder.calls[0][0]
    target_text = "\n".join(chunk.text for chunk in vectors.versions[accepted.version_id])
    assert "6500" in target_text
    assert "6800" not in target_text
    assert db.get_current_document_version(document_id)["id"] == accepted.version_id
    assert db.get_structured_dataset_for_version(accepted.version_id)["status"] == "ready"
    assert db.get_spreadsheet_edit_session(accepted.session_id)["status"] == "active"
    hits = db.search_material_search_fts("6500")
    assert hits[0]["version_id"] == accepted.version_id

    structure = SpreadsheetEditService(
        db, versions_root, datasets_root
    ).save_structure(
        document_id,
        owner,
        SpreadsheetStructureEditRequest(
            base_version_id=accepted.version_id,
            action="insert_rows",
            sheet_key="sheet_0001",
            index=2,
        ),
        RequestMeta(),
    )
    third = process_index_job(
        {
            "id": 3,
            "document_id": document_id,
            "version_id": structure.version_id,
        },
        services,
    )
    assert third.status == "done"
    assert db.get_current_document_version(document_id)["id"] == structure.version_id
    assert db.get_structured_dataset_for_version(structure.version_id)["status"] == "ready"
    assert db.get_spreadsheet_edit_session(structure.session_id)["status"] == "active"


def test_index_job_publishes_the_schema_mirror_without_a_restart(tmp_path):
    """A worker-created dataset must be visible to the catalog immediately.

    The authorized catalog reads the ``users.db`` mirror, so indexing a version
    without writing that mirror made the new dataset invisible until the next
    startup backfill.
    """
    db = UserDB(str(tmp_path / "users.sqlite3"))
    db.create_material_schema()
    owner_id = db.create_user("owner", "Password-123!")
    document_id = db.create_document_permission("two-sheets.xlsx", "rag_docs", owner_id)
    versions_root = tmp_path / "versions"
    source = final_version_path(versions_root, document_id, 1, "two-sheets.xlsx")
    source.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    lens = workbook.active
    lens.title = "镜头"
    lens.append(["品牌", "售价"])
    lens.append(["腾龙", 320])
    body = workbook.create_sheet("机身")
    body.append(["品牌", "单价"])
    body.append(["佳能", 399])
    workbook.save(source)
    version_id = db.create_document_version(
        document_id, "two-sheets.xlsx", str(source), "hash-mirror", owner_id
    )
    services = WorkerServices(
        db=db,
        versions_root=versions_root,
        structured_datasets_root=tmp_path / "datasets",
        registry=build_default_format_registry(),
        vector_store=MemoryVectorStore(),
        embedder=RecordingEmbedder(),
    )

    result = process_index_job(
        {"id": 1, "document_id": document_id, "version_id": version_id}, services
    )

    assert result.status == "done"
    dataset = db.get_structured_dataset_for_version(version_id)
    assert dataset["status"] == "ready"
    rows = db.list_structured_dataset_schema(int(dataset["id"]))
    assert {row["sheet_key"] for row in rows} == {"sheet_0001", "sheet_0002"}
    # Both sheets own a c_0002, which the mirror has to keep distinct.
    assert sorted(row["column_key"] for row in rows) == ["c_0001", "c_0001", "c_0002", "c_0002"]
    assert len({(row["sheet_key"], row["column_key"]) for row in rows}) == 4
    db.close()


def test_real_vector_copy_rebinds_unchanged_unit_to_target_dataset(
    tmp_path, monkeypatch
):
    from qdrant_client.models import FieldCondition, Filter, MatchValue

    import rag.vector_store as vector_store
    from config import settings

    monkeypatch.setattr(settings, "qdrant_path", str(tmp_path / "qdrant"))
    monkeypatch.setattr(settings, "embed_dimension", 4)
    vector_store._ensure_collection_name("sheet-kb")
    source = Chunk(
        text="unchanged",
        doc_name="prices.xlsx",
        chunk_index=2,
        metadata={"unit_key": "sheet_0001:rows:2:2"},
        document_id=12,
        document_version_id=7,
        version_no=1,
        dataset_id=70,
        sheet_key="sheet_0001",
        chunk_kind="row_group",
        row_start=2,
        row_end=2,
    )
    vector_store.add_to_collection(
        "sheet-kb",
        [source],
        [[0.1, 0.2, 0.3, 0.4]],
        document_id=12,
        document_version_id=7,
        version_no=1,
    )

    copied = vector_store.copy_units_to_version(
        "sheet-kb",
        source_version_id=7,
        target_version_id=8,
        document_id=12,
        version_no=2,
        unit_keys=["sheet_0001:rows:2:2"],
        dataset_id=80,
    )

    points, _ = vector_store._get_client().scroll(
        collection_name="sheet-kb",
        scroll_filter=Filter(
            must=[
                FieldCondition(
                    key="document_version_id", match=MatchValue(value=8)
                )
            ]
        ),
        with_payload=True,
        with_vectors=True,
    )
    assert copied == 1
    assert len(points) == 1
    assert points[0].payload["dataset_id"] == 80
    assert points[0].payload["document_version_id"] == 8
