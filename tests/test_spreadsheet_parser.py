from datetime import date

import pytest
from openpyxl import Workbook

from config import settings
from rag.spreadsheets.parser import SpreadsheetLimitError, parse_spreadsheet


def _write_workbook(tmp_path):
    path = tmp_path / "sales.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "Sales"
    sheet.append(["zero", "flag", "missing", "amount", "when", "formula"])
    sheet.append([0, False, None, "12.50", date(2026, 8, 10), "=A2+1"])
    hidden = book.create_sheet("Archive")
    hidden.sheet_state = "hidden"
    hidden.append(["name"])
    hidden.append(["old"])
    book.save(path)
    return path


def test_xlsx_preserves_zero_false_empty_and_formula_text(tmp_path):
    parsed = parse_spreadsheet(_write_workbook(tmp_path), "xlsx", None)

    row = parsed.sheets[0].rows[0]
    assert row.cells[0].value == 0
    assert row.cells[1].value is False
    assert row.cells[2].value is None
    assert parsed.sheets[0].formula_cells[0].formula == "=A2+1"
    assert parsed.sheets[0].formula_cells[0].cached_value is None
    assert parsed.sheets[1].visibility == "hidden"


@pytest.mark.parametrize("encoding", ["utf-8-sig", "utf-8", "gb18030"])
def test_csv_encodings_and_delimiters_round_trip(tmp_path, encoding):
    path = tmp_path / f"sales-{encoding}.csv"
    path.write_bytes("姓名;金额\n张三;0\n".encode(encoding))

    parsed = parse_spreadsheet(path, "csv", None)

    assert parsed.sheets[0].headers == ("姓名", "金额")
    assert parsed.sheets[0].rows[0].cells[1].value == 0


def test_merged_or_duplicate_headers_require_explicit_mapping(tmp_path):
    path = tmp_path / "needs-mapping.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.merge_cells("A1:B1")
    sheet["A1"] = "Sales"
    sheet.append(["region", "amount"])
    sheet.append(["east", "1"])
    duplicate = book.create_sheet("Duplicate")
    duplicate.append(["name", "name"])
    duplicate.append(["east", "west"])
    book.save(path)

    parsed = parse_spreadsheet(path, "xlsx", None)

    assert parsed.needs_mapping is True
    assert {"merged_header", "duplicate_header"} <= set(parsed.issues)


def test_parser_enforces_configured_row_limit(tmp_path, monkeypatch):
    path = _write_workbook(tmp_path)
    monkeypatch.setattr(settings, "spreadsheet_max_rows", 0)

    with pytest.raises(SpreadsheetLimitError, match="row limit"):
        parse_spreadsheet(path, "xlsx", None)
