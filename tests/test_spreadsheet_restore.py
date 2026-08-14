import hashlib
import json

import pytest
from fastapi import HTTPException
from openpyxl import Workbook, load_workbook

from rag.document_version_store import final_version_path
from rag.spreadsheets.edit_models import RequestMeta, SpreadsheetRestoreRequest
from rag.user_db import UserDB, VersionConflict


def _add_workbook_version(db, root, document_id, owner_id, version_no, value):
    path = final_version_path(root, document_id, version_no, "prices.xlsx")
    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    book.active.title = "Prices"
    book.active["A1"] = value
    book.save(path)
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    version_id = db.create_document_version(
        document_id, "prices.xlsx", str(path), digest, owner_id
    )
    db.activate_document_version(
        document_id, version_id, "rag_docs", "prices.xlsx"
    )
    return version_id, path


@pytest.fixture
def restore_context(tmp_path):
    db = UserDB(str(tmp_path / "users.db"))
    db.create_material_schema()
    owner_id = db.create_user("owner", "Password-123!")
    outsider_id = db.create_user("outsider", "Password-123!")
    document_id = db.create_document_permission(
        "prices.xlsx", "rag_docs", owner_id
    )
    root = tmp_path / "document_versions"
    versions = [
        _add_workbook_version(db, root, document_id, owner_id, number, value)
        for number, value in ((1, "v1"), (2, "v2"), (3, "v3"))
    ]
    other_document_id = db.create_document_permission(
        "other.xlsx", "rag_docs", owner_id
    )
    other_version_id, _ = _add_workbook_version(
        db, root, other_document_id, owner_id, 1, "other"
    )
    yield {
        "db": db,
        "owner": db.get_user_by_id(owner_id),
        "outsider": db.get_user_by_id(outsider_id),
        "document_id": document_id,
        "versions": [item[0] for item in versions],
        "paths": [item[1] for item in versions],
        "other_version_id": other_version_id,
        "root": root,
        "structured_root": tmp_path / "structured",
    }
    db.close()


def _service(context, enqueue=None):
    from rag.spreadsheets.edit_service import SpreadsheetEditService

    return SpreadsheetEditService(
        context["db"],
        context["root"],
        context["structured_root"],
        enqueue=enqueue,
    )


def test_restore_creates_a_new_version_and_retains_all_history(restore_context):
    context = restore_context
    v1, v2, v3 = context["versions"]
    jobs = []

    def enqueue(db, document_id, version_id, content_hash, not_before):
        jobs.append(version_id)
        return db.enqueue_latest_index_job(
            document_id, version_id, content_hash, not_before
        )

    accepted = _service(context, enqueue).restore(
        context["document_id"],
        v1,
        context["owner"],
        SpreadsheetRestoreRequest(base_version_id=v3, reason="undo changes"),
        RequestMeta(),
    )

    assert accepted.status == "pending"
    assert jobs == [accepted.version_id]
    versions = context["db"].list_document_versions(context["document_id"])
    assert [item["version_no"] for item in versions] == [4, 3, 2, 1]
    assert context["db"].get_current_document_version(
        context["document_id"]
    )["id"] == v3
    restored = context["db"].get_document_version_by_id(accepted.version_id)
    assert restored["restored_from_version_id"] == v1
    assert load_workbook(restored["file_path"])["Prices"]["A1"].value == "v1"
    session = context["db"].get_spreadsheet_edit_session(accepted.session_id)
    assert session["operation"] == "rollback"
    assert session["change_count"] == 0
    assert json.loads(session["risk_flags_json"]) == ["non_previous_restore"]

    context["db"].activate_document_version(
        context["document_id"], accepted.version_id, "rag_docs", "prices.xlsx"
    )
    second = _service(context, enqueue).restore(
        context["document_id"],
        v3,
        context["owner"],
        SpreadsheetRestoreRequest(base_version_id=accepted.version_id),
        RequestMeta(),
    )
    assert context["db"].get_document_version_by_id(second.version_id)[
        "version_no"
    ] == 5
    assert len(context["db"].list_document_versions(context["document_id"])) == 5


def test_restore_rejects_permission_stale_base_and_cross_document_source(
    restore_context,
):
    context = restore_context
    v1, _, v3 = context["versions"]
    service = _service(context)

    with pytest.raises(HTTPException) as denied:
        service.restore(
            context["document_id"],
            v1,
            context["outsider"],
            SpreadsheetRestoreRequest(base_version_id=v3),
            RequestMeta(),
        )
    assert denied.value.status_code == 403
    with pytest.raises(VersionConflict, match="spreadsheet_base_version_conflict"):
        service.restore(
            context["document_id"],
            v1,
            context["owner"],
            SpreadsheetRestoreRequest(base_version_id=99999),
            RequestMeta(),
        )
    with pytest.raises(VersionConflict, match="spreadsheet_restore_source_conflict"):
        service.restore(
            context["document_id"],
            context["other_version_id"],
            context["owner"],
            SpreadsheetRestoreRequest(base_version_id=v3),
            RequestMeta(),
        )
    assert len(context["db"].list_document_versions(context["document_id"])) == 3


def test_restore_missing_source_file_creates_nothing(restore_context):
    from rag.spreadsheets.edit_service import SpreadsheetEditServiceError

    context = restore_context
    v1, _, v3 = context["versions"]
    context["paths"][0].unlink()

    with pytest.raises(SpreadsheetEditServiceError, match="spreadsheet_version_missing"):
        _service(context).restore(
            context["document_id"],
            v1,
            context["owner"],
            SpreadsheetRestoreRequest(base_version_id=v3),
            RequestMeta(),
        )

    assert len(context["db"].list_document_versions(context["document_id"])) == 3
    assert context["db"]._conn.execute(
        "SELECT COUNT(*) FROM spreadsheet_edit_sessions"
    ).fetchone()[0] == 0


def test_restore_endpoint_returns_pending_and_serializes_provenance(
    isolated_api, make_user
):
    owner, headers = make_user("restore-owner")
    db = isolated_api.app.state.user_db
    document_id = db.create_document_permission(
        "prices.xlsx", "rag_docs", owner["id"]
    )
    v1, _ = _add_workbook_version(
        db,
        isolated_api.app.state.document_versions_root,
        document_id,
        owner["id"],
        1,
        "v1",
    )
    v2, _ = _add_workbook_version(
        db,
        isolated_api.app.state.document_versions_root,
        document_id,
        owner["id"],
        2,
        "v2",
    )

    response = isolated_api.post(
        f"/documents/{document_id}/versions/{v1}/restore",
        headers=headers,
        json={"base_version_id": v2, "reason": "undo"},
    )

    assert response.status_code == 202
    body = response.json()
    versions = isolated_api.get(
        f"/documents/{document_id}/versions", headers=headers
    )
    assert versions.status_code == 200
    restored = next(
        item for item in versions.json()["versions"]
        if item["id"] == body["version_id"]
    )
    assert restored["restored_from_version_id"] == v1
    assert "file_path" not in versions.text + response.text
