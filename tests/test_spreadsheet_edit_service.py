import hashlib
from pathlib import Path

import pytest
from fastapi import HTTPException
from openpyxl import Workbook, load_workbook

from rag.document_version_store import final_version_path
from rag.spreadsheets.edit_models import (
    CellPatch,
    RequestMeta,
    SpreadsheetEditRequest,
    SpreadsheetStructureEditRequest,
)
from rag.spreadsheets.parser import parse_spreadsheet
from rag.spreadsheets.store import resolve_dataset_path, write_dataset
from rag.user_db import UserDB, VersionConflict


@pytest.fixture
def spreadsheet_service_context(tmp_path):
    db = UserDB(str(tmp_path / "users.db"))
    db.create_material_schema()
    owner_id = db.create_user("owner", "Password-123!")
    outsider_id = db.create_user("outsider", "Password-123!")
    document_id = db.create_document_permission(
        "prices.xlsx", "rag_docs", owner_id
    )
    versions_root = tmp_path / "document_versions"
    structured_root = tmp_path / "structured_datasets"
    source = final_version_path(
        versions_root, document_id, 1, "prices.xlsx"
    )
    source.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    book.active.title = "Prices"
    book.active.append(["product", "price"])
    book.active.append(["Phone", 6800])
    book.save(source)
    source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    version_id = db.create_document_version(
        document_id, "prices.xlsx", str(source), source_hash, owner_id
    )
    dataset_id = db.create_structured_dataset(
        document_id, version_id, "xlsx", "row_groups"
    )
    dataset_path = resolve_dataset_path(
        structured_root, document_id, version_id
    )
    write_dataset(dataset_path, parse_spreadsheet(source, "xlsx", None))
    db.activate_document_and_dataset(
        document_id=document_id,
        version_id=version_id,
        dataset_id=dataset_id,
        kb_id="rag_docs",
        filename="prices.xlsx",
    )
    owner = db.get_user_by_id(owner_id)
    outsider = db.get_user_by_id(outsider_id)
    yield {
        "db": db,
        "owner": owner,
        "outsider": outsider,
        "document_id": document_id,
        "version_id": version_id,
        "versions_root": versions_root,
        "structured_root": structured_root,
    }
    db.close()


def _service(context, enqueue=None):
    from rag.spreadsheets.edit_service import SpreadsheetEditService

    return SpreadsheetEditService(
        context["db"],
        context["versions_root"],
        context["structured_root"],
        enqueue=enqueue,
    )


def _request(base_version_id, value=6500):
    return SpreadsheetEditRequest(
        base_version_id=base_version_id,
        changes=(
            CellPatch(
                sheet_key="sheet_0001",
                cell="B2",
                value_type="number",
                value=value,
            ),
        ),
        reason="price update",
    )


def test_save_creates_one_pending_version_audit_and_index_job(
    spreadsheet_service_context
):
    context = spreadsheet_service_context
    enqueued = []

    def record_enqueue(db, document_id, version_id, content_hash, not_before):
        enqueued.append((document_id, version_id, content_hash, not_before))
        return db.enqueue_latest_index_job(
            document_id, version_id, content_hash, not_before
        )

    result = _service(context, record_enqueue).save(
        context["document_id"],
        context["owner"],
        _request(context["version_id"]),
        RequestMeta(ip_address="127.0.0.1", user_agent="pytest"),
    )

    assert result.status == "pending"
    assert len(enqueued) == 1
    assert enqueued[0][0] == context["document_id"]
    assert enqueued[0][1] == result.version_id
    assert context["db"].get_current_document_version(
        context["document_id"]
    )["id"] == context["version_id"]
    created = context["db"].get_document_version_by_id(result.version_id)
    assert created["status"] == "staging"
    assert Path(created["file_path"]).is_file()
    assert load_workbook(created["file_path"])["Prices"]["B2"].value == 6500
    session = context["db"].get_spreadsheet_edit_session(result.session_id)
    assert session["new_version_id"] == result.version_id
    stored = context["db"]._conn.execute(
        "SELECT old_value_json, new_value_json FROM spreadsheet_cell_changes "
        "WHERE session_id = ?",
        (result.session_id,),
    ).fetchone()
    assert tuple(stored) == ("6800", "6500")


def test_unchanged_edit_creates_no_version_session_or_job(
    spreadsheet_service_context
):
    context = spreadsheet_service_context
    jobs = []
    service = _service(context, lambda *args: jobs.append(args))

    result = service.save(
        context["document_id"],
        context["owner"],
        _request(context["version_id"], value=6800),
        RequestMeta(),
    )

    assert result.status == "unchanged"
    assert len(context["db"].list_document_versions(context["document_id"])) == 1
    assert context["db"]._conn.execute(
        "SELECT COUNT(*) FROM spreadsheet_edit_sessions"
    ).fetchone()[0] == 0
    assert jobs == []


def test_save_rejects_stale_base_and_non_editor(spreadsheet_service_context):
    context = spreadsheet_service_context
    service = _service(context)

    with pytest.raises(VersionConflict, match="spreadsheet_base_version_conflict"):
        service.save(
            context["document_id"],
            context["owner"],
            _request(99999),
            RequestMeta(),
        )
    with pytest.raises(HTTPException) as denied:
        service.save(
            context["document_id"],
            context["outsider"],
            _request(context["version_id"]),
            RequestMeta(),
        )
    assert denied.value.status_code == 403
    assert len(context["db"].list_document_versions(context["document_id"])) == 1


def test_save_uses_server_sheet_map_and_rejects_unknown_sheet(
    spreadsheet_service_context
):
    from rag.spreadsheets.edit_service import SpreadsheetEditServiceError

    context = spreadsheet_service_context
    request = SpreadsheetEditRequest(
        base_version_id=context["version_id"],
        changes=(
            CellPatch(
                sheet_key="invented_by_client",
                cell="B2",
                value_type="number",
                value=6500,
            ),
        ),
    )

    with pytest.raises(SpreadsheetEditServiceError, match="unknown_sheet"):
        _service(context).save(
            context["document_id"], context["owner"], request, RequestMeta()
        )


def test_rejected_formula_edit_is_value_free_audit(spreadsheet_service_context):
    from rag.spreadsheets.edit_service import SpreadsheetEditServiceError

    context = spreadsheet_service_context
    current = context["db"].get_current_document_version(context["document_id"])
    workbook = load_workbook(current["file_path"])
    workbook["Prices"]["C2"] = "=1+1"
    workbook.save(current["file_path"])
    request = SpreadsheetEditRequest(
        base_version_id=context["version_id"],
        changes=(
            CellPatch(
                sheet_key="sheet_0001",
                cell="C2",
                value_type="number",
                value=2,
            ),
        ),
    )

    with pytest.raises(
        SpreadsheetEditServiceError, match="formula_cell_not_editable"
    ):
        _service(context).save(
            context["document_id"], context["owner"], request, RequestMeta()
        )

    session = context["db"]._conn.execute(
        "SELECT status, new_version_id, change_count, risk_flags_json "
        "FROM spreadsheet_edit_sessions ORDER BY id DESC LIMIT 1"
    ).fetchone()
    assert tuple(session) == (
        "rejected",
        None,
        0,
        '["rejected_formula_edit"]',
    )


def test_structure_save_creates_pending_version_and_structure_audit(
    spreadsheet_service_context,
):
    context = spreadsheet_service_context
    enqueued = []

    result = _service(context, lambda *args: enqueued.append(args)).save_structure(
        context["document_id"],
        context["owner"],
        SpreadsheetStructureEditRequest(
            base_version_id=context["version_id"],
            action="insert_rows",
            sheet_key="sheet_0001",
            index=2,
        ),
        RequestMeta(),
    )

    assert result.status == "pending"
    assert len(enqueued) == 1
    created = context["db"].get_document_version_by_id(result.version_id)
    workbook = load_workbook(created["file_path"])
    assert workbook["Prices"]["A2"].value is None
    assert workbook["Prices"]["A3"].value == "Phone"
    stored = context["db"]._conn.execute(
        "SELECT action, sheet_key, target_index, count "
        "FROM spreadsheet_structure_changes WHERE session_id=?",
        (result.session_id,),
    ).fetchone()
    assert tuple(stored) == ("insert_rows", "sheet_0001", 2, 1)


def test_structure_save_records_rejected_nonempty_delete_without_values(
    spreadsheet_service_context,
):
    from rag.spreadsheets.edit_service import SpreadsheetEditServiceError

    context = spreadsheet_service_context
    request = SpreadsheetStructureEditRequest(
        base_version_id=context["version_id"],
        action="delete_rows",
        sheet_key="sheet_0001",
        index=2,
    )

    with pytest.raises(
        SpreadsheetEditServiceError,
        match="nonempty_delete_requires_confirmation",
    ):
        _service(context).save_structure(
            context["document_id"], context["owner"], request, RequestMeta()
        )

    session = context["db"]._conn.execute(
        "SELECT status, new_version_id, change_count, error_code "
        "FROM spreadsheet_edit_sessions ORDER BY id DESC LIMIT 1"
    ).fetchone()
    assert tuple(session) == (
        "rejected",
        None,
        0,
        "nonempty_delete_requires_confirmation",
    )
