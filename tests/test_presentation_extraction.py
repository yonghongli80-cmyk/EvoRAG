"""Task 3: page-aware PPTX extraction without image understanding."""

import base64
from pathlib import Path

import pytest

from rag.materials.presentation import extract_presentation

MINIMAL_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
)


@pytest.fixture
def pptx_fixture(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    path = tmp_path / "policy.pptx"

    def save_with(build) -> Path:
        prs = Presentation()
        build(prs)
        prs.save(str(path))
        return path

    class Fixture:
        def __init__(self):
            self.path = self._build_default()

        def _build_default(self) -> Path:
            def build(prs):
                slide1 = prs.slides.add_slide(prs.slide_layouts[5])
                slide1.shapes.title.text = "年度制度"
                slide1.notes_slide.notes_text_frame.text = "讲者备注"
                slide2 = prs.slides.add_slide(prs.slide_layouts[5])
                table = slide2.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(1))
                table.table.cell(0, 0).text = "部门"
                table.table.cell(0, 1).text = "负责人"
                table.table.cell(1, 0).text = "市场部"
                table.table.cell(1, 1).text = "张三"

            return save_with(build)

        def with_image_only_slide(self) -> Path:
            def build(prs):
                img = tmp_path / "img.png"
                img.write_bytes(MINIMAL_PNG)
                slide = prs.slides.add_slide(prs.slide_layouts[5])
                slide.shapes.add_picture(str(img), Inches(1), Inches(1))

            self.path = save_with(build)
            return self.path

    return Fixture()


def test_pptx_extracts_slide_text_notes_and_table_with_page_citations(pptx_fixture):
    result = extract_presentation(pptx_fixture.path, version_id=8)
    assert [(unit.location.kind, unit.location.value) for unit in result.units] == [
        ("slide", "1"),
        ("slide", "2"),
    ]
    assert "年度制度" in result.units[0].text
    assert "讲者备注" in result.units[0].text
    assert "部门 | 负责人" in result.units[1].text


def test_pptx_does_not_ocr_or_caption_embedded_image(pptx_fixture):
    class ForbiddenVisionClient:
        def __init__(self):
            self.calls = []

        def caption(self, *_args, **_kwargs):
            self.calls.append(1)

    forbidden_vision_client = ForbiddenVisionClient()

    result = extract_presentation(pptx_fixture.with_image_only_slide(), version_id=8)

    assert forbidden_vision_client.calls == []
    assert result.units[-1].text == ""
