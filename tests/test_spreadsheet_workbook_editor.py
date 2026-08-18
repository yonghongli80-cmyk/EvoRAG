from datetime import date

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from rag.spreadsheets.edit_models import CellPatch


def _write_workbook(path):
    book = Workbook()
    prices = book.active
    prices.title = "Prices"
    prices.freeze_panes = "A2"
    prices.merge_cells("D1:E1")
    prices["A1"] = "Product"
    prices["A1"].font = Font(bold=True)
    prices["A1"].fill = PatternFill("solid", fgColor="FFFF00")
    prices["D1"] = "Merged heading"
    prices["A2"] = "Phone"
    prices["B2"] = 6800
    prices["B2"].number_format = "0.00"
    prices["C2"] = "=B2*0.9"
    prices["F2"] = True
    prices["G2"] = date(2026, 8, 1)
    prices["H2"] = "remove me"
    prices.row_dimensions[2].height = 24
    prices.column_dimensions["B"].width = 18

    archive = book.create_sheet("Archive")
    archive.sheet_state = "hidden"
    archive["A2"] = "old"
    book.security.lockStructure = True
    book.save(path)


def test_apply_cell_patches_preserves_workbook_structure_and_formatting(tmp_path):
    from rag.spreadsheets.workbook_editor import apply_cell_patches

    source = tmp_path / "source.xlsx"
    destination = tmp_path / "edited.xlsx"
    _write_workbook(source)
    before = load_workbook(source, data_only=False)
    original_style = before["Prices"]["B2"].style_id

    changes = apply_cell_patches(
        source,
        destination,
        {"sheet_0001": "Prices", "sheet_0002": "Archive"},
        (
            CellPatch(sheet_key="sheet_0001", cell="A2", value_type="text", value="Tablet"),
            CellPatch(sheet_key="sheet_0001", cell="B2", value_type="number", value=6500),
            CellPatch(sheet_key="sheet_0001", cell="F2", value_type="boolean", value=False),
            CellPatch(sheet_key="sheet_0001", cell="G2", value_type="date", value="2026-08-14"),
            CellPatch(sheet_key="sheet_0001", cell="H2", value_type="blank", value=None),
            CellPatch(sheet_key="sheet_0002", cell="A2", value_type="text", value="new"),
        ),
    )

    edited = load_workbook(destination, data_only=False)
    assert edited.sheetnames == ["Prices", "Archive"]
    assert edited["Prices"]["A2"].value == "Tablet"
    assert edited["Prices"]["B2"].value == 6500
    assert edited["Prices"]["F2"].value is False
    assert edited["Prices"]["G2"].value.date() == date(2026, 8, 14)
    assert edited["Prices"]["H2"].value is None
    assert edited["Archive"]["A2"].value == "new"
    assert edited["Prices"]["C2"].value == "=B2*0.9"
    assert edited["Prices"]["B2"].style_id == original_style
    assert [str(value) for value in edited["Prices"].merged_cells.ranges] == ["D1:E1"]
    assert edited["Archive"].sheet_state == "hidden"
    assert edited["Prices"].freeze_panes == "A2"
    assert edited["Prices"].row_dimensions[2].height == 24
    assert edited["Prices"].column_dimensions["B"].width == 18
    assert edited.security.lockStructure is True
    assert [change.cell_address for change in changes] == [
        "A2", "B2", "F2", "G2", "H2", "A2"
    ]
    assert changes[1].old_value == 6800
    assert changes[1].new_value == 6500


def test_apply_cell_patches_supports_workbooks_without_protection_metadata(tmp_path):
    from rag.spreadsheets.workbook_editor import apply_cell_patches

    source = tmp_path / "unprotected.xlsx"
    destination = tmp_path / "edited.xlsx"
    book = Workbook()
    book.active.title = "Sheet1"
    book.active["A1"] = "before"
    book.security = None
    book.save(source)

    changes = apply_cell_patches(
        source,
        destination,
        {"sheet_0001": "Sheet1"},
        (
            CellPatch(
                sheet_key="sheet_0001",
                cell="A1",
                value_type="text",
                value="after",
            ),
        ),
    )

    edited = load_workbook(destination, data_only=False)
    assert edited["Sheet1"]["A1"].value == "after"
    assert edited.security is None
    assert changes[0].old_value == "before"
    assert changes[0].new_value == "after"


def _assert_edit_error(source, destination, patches, reason_code, sheet_map=None):
    from rag.spreadsheets.workbook_editor import WorkbookEditError, apply_cell_patches

    with pytest.raises(WorkbookEditError) as caught:
        apply_cell_patches(
            source,
            destination,
            sheet_map or {"sheet_0001": "Prices"},
            patches,
        )
    assert caught.value.reason_code == reason_code
    assert str(caught.value) == reason_code
    assert not destination.exists()


@pytest.mark.parametrize(
    ("cell", "reason_code"),
    [
        ("C2", "formula_cell_not_editable"),
        ("E1", "merged_cell_not_editable"),
        ("A0", "invalid_cell_address"),
        ("A1:B2", "invalid_cell_address"),
    ],
)
def test_rejects_formula_merged_and_invalid_targets(
    tmp_path, cell, reason_code
):
    source = tmp_path / "source.xlsx"
    destination = tmp_path / "edited.xlsx"
    _write_workbook(source)

    _assert_edit_error(
        source,
        destination,
        (CellPatch(sheet_key="sheet_0001", cell=cell, value_type="text", value="safe"),),
        reason_code,
    )


def test_rejects_locked_cells_on_a_protected_sheet(tmp_path):
    source = tmp_path / "source.xlsx"
    destination = tmp_path / "edited.xlsx"
    _write_workbook(source)
    book = load_workbook(source)
    book["Prices"].protection.sheet = True
    book.save(source)

    _assert_edit_error(
        source,
        destination,
        (CellPatch(sheet_key="sheet_0001", cell="B2", value_type="number", value=6500),),
        "protected_cell",
    )


def test_rejects_duplicate_cells_and_more_than_500_patches(tmp_path):
    source = tmp_path / "source.xlsx"
    _write_workbook(source)
    duplicate = CellPatch(
        sheet_key="sheet_0001", cell="B2", value_type="number", value=6500
    )
    _assert_edit_error(
        source,
        tmp_path / "duplicate.xlsx",
        (duplicate, duplicate),
        "duplicate_cell_patch",
    )
    too_many = tuple(
        CellPatch(
            sheet_key="sheet_0001",
            cell=f"A{row}",
            value_type="number",
            value=row,
        )
        for row in range(1, 502)
    )
    _assert_edit_error(
        source,
        tmp_path / "too-many.xlsx",
        too_many,
        "invalid_patch_count",
    )


def test_rejects_unknown_sheets_and_formula_injection(tmp_path):
    source = tmp_path / "source.xlsx"
    _write_workbook(source)
    _assert_edit_error(
        source,
        tmp_path / "unknown.xlsx",
        (CellPatch(sheet_key="missing", cell="A1", value_type="text", value="safe"),),
        "unknown_sheet",
    )
    _assert_edit_error(
        source,
        tmp_path / "formula-injection.xlsx",
        (CellPatch(sheet_key="sheet_0001", cell="A2", value_type="text", value="=1+1"),),
        "formula_injection",
    )


@pytest.mark.parametrize("value", [float("nan"), float("inf"), float("-inf")])
def test_editor_rejects_non_finite_numbers_even_if_model_validation_is_bypassed(
    tmp_path, value
):
    source = tmp_path / "source.xlsx"
    _write_workbook(source)
    unsafe_patch = CellPatch.model_construct(
        sheet_key="sheet_0001", cell="B2", value_type="number", value=value
    )

    _assert_edit_error(
        source,
        tmp_path / "edited.xlsx",
        (unsafe_patch,),
        "invalid_cell_value",
    )


def test_rejects_macro_enabled_workbooks(tmp_path):
    source = tmp_path / "source.xlsm"
    destination = tmp_path / "edited.xlsx"
    _write_workbook(source)

    _assert_edit_error(
        source,
        destination,
        (CellPatch(sheet_key="sheet_0001", cell="B2", value_type="number", value=6500),),
        "unsupported_workbook_format",
    )


def test_post_save_drift_deletes_the_destination(tmp_path, monkeypatch):
    import rag.spreadsheets.workbook_editor as editor

    source = tmp_path / "source.xlsx"
    destination = tmp_path / "edited.xlsx"
    _write_workbook(source)
    real_snapshot = editor._snapshot_workbook
    calls = 0

    def drifting_snapshot(book):
        nonlocal calls
        calls += 1
        snapshot = real_snapshot(book)
        if calls == 2:
            snapshot["sheet_names"] = ("tampered",)
        return snapshot

    monkeypatch.setattr(editor, "_snapshot_workbook", drifting_snapshot)

    _assert_edit_error(
        source,
        destination,
        (CellPatch(sheet_key="sheet_0001", cell="B2", value_type="number", value=6500),),
        "workbook_validation_failed",
    )
