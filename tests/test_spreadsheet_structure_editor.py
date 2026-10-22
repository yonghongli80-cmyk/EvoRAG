from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from rag.spreadsheets.edit_models import SpreadsheetStructureEditRequest


def _save_workbook(path: Path) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "Data"
    sheet["A1"] = "header"
    sheet["A2"] = "first"
    sheet["B2"] = "second"
    sheet["A3"] = "third"
    sheet["A2"].fill = PatternFill(fill_type="solid", fgColor="FFFF00")
    book.create_sheet("Other")["A1"] = "untouched"
    book.save(path)


def _apply(source: Path, destination: Path, **values):
    from rag.spreadsheets.structure_editor import apply_structure_edit

    request = SpreadsheetStructureEditRequest(base_version_id=1, **values)
    return apply_structure_edit(
        source,
        destination,
        {"sheet_0001": "Data", "sheet_0002": "Other"},
        request,
    )


def test_insert_rows_shifts_values_and_styles_without_touching_other_sheet(tmp_path):
    source = tmp_path / "source.xlsx"
    destination = tmp_path / "edited.xlsx"
    _save_workbook(source)

    change = _apply(
        source,
        destination,
        action="insert_rows",
        sheet_key="sheet_0001",
        index=2,
        count=1,
    )

    book = load_workbook(destination)
    assert book["Data"]["A2"].value is None
    assert book["Data"]["A3"].value == "first"
    assert book["Data"]["A3"].fill.fgColor.rgb == "00FFFF00"
    assert book["Other"]["A1"].value == "untouched"
    assert change.action == "insert_rows"
    assert change.affected_nonempty_cells == 0


def test_insert_columns_shifts_existing_cells(tmp_path):
    source = tmp_path / "source.xlsx"
    destination = tmp_path / "edited.xlsx"
    _save_workbook(source)

    _apply(
        source,
        destination,
        action="insert_columns",
        sheet_key="sheet_0001",
        index=2,
        count=2,
    )

    book = load_workbook(destination)
    assert book["Data"]["B2"].value is None
    assert book["Data"]["D2"].value == "second"


def test_insert_rows_accepts_workbook_without_security_metadata(tmp_path):
    source = tmp_path / "source.xlsx"
    destination = tmp_path / "edited.xlsx"
    _save_workbook(source)

    book = load_workbook(source)
    book.security = None
    book.save(source)
    assert load_workbook(source).security is None

    _apply(
        source,
        destination,
        action="insert_rows",
        sheet_key="sheet_0001",
        index=1,
        count=1,
    )

    assert load_workbook(destination)["Data"]["A2"].value == "header"


@pytest.mark.parametrize(
    ("action", "index", "count", "expected_count"),
    [
        ("delete_rows", 2, 1, 2),
        ("delete_columns", 1, 1, 3),
    ],
)
def test_nonempty_delete_requires_confirmation(
    tmp_path, action, index, count, expected_count
):
    from rag.spreadsheets.structure_editor import WorkbookStructureError

    source = tmp_path / "source.xlsx"
    destination = tmp_path / "edited.xlsx"
    _save_workbook(source)

    with pytest.raises(WorkbookStructureError) as caught:
        _apply(
            source,
            destination,
            action=action,
            sheet_key="sheet_0001",
            index=index,
            count=count,
        )
    assert caught.value.reason_code == "nonempty_delete_requires_confirmation"
    assert not destination.exists()

    change = _apply(
        source,
        destination,
        action=action,
        sheet_key="sheet_0001",
        index=index,
        count=count,
        confirm_nonempty=True,
    )
    assert change.affected_nonempty_cells == expected_count


def test_confirmed_delete_rows_shifts_following_values(tmp_path):
    source = tmp_path / "source.xlsx"
    destination = tmp_path / "edited.xlsx"
    _save_workbook(source)

    _apply(
        source,
        destination,
        action="delete_rows",
        sheet_key="sheet_0001",
        index=2,
        confirm_nonempty=True,
    )

    assert load_workbook(destination)["Data"]["A2"].value == "third"


def test_add_sheet_appends_a_blank_named_sheet(tmp_path):
    source = tmp_path / "source.xlsx"
    destination = tmp_path / "edited.xlsx"
    _save_workbook(source)

    change = _apply(
        source,
        destination,
        action="add_sheet",
        new_sheet_name="August",
    )

    book = load_workbook(destination)
    assert book.sheetnames == ["Data", "Other", "August"]
    assert book["August"].max_row == 1
    assert book["August"]["A1"].value is None
    assert change.new_sheet_name == "August"


@pytest.mark.parametrize("new_name", ["data", "bad/name"])
def test_add_sheet_rejects_duplicate_or_illegal_names(tmp_path, new_name):
    from rag.spreadsheets.structure_editor import WorkbookStructureError

    source = tmp_path / "source.xlsx"
    destination = tmp_path / "edited.xlsx"
    _save_workbook(source)

    with pytest.raises(WorkbookStructureError) as caught:
        _apply(source, destination, action="add_sheet", new_sheet_name=new_name)
    assert caught.value.reason_code == "invalid_sheet_name"
    assert not destination.exists()


@pytest.mark.parametrize("unsafe_kind", ["formula", "merge", "protection"])
def test_row_operations_refuse_unsafe_workbooks(tmp_path, unsafe_kind):
    from rag.spreadsheets.structure_editor import WorkbookStructureError

    source = tmp_path / "source.xlsx"
    destination = tmp_path / "edited.xlsx"
    _save_workbook(source)
    book = load_workbook(source)
    if unsafe_kind == "formula":
        book["Other"]["B1"] = "=1+1"
        expected = "workbook_contains_formulas"
    elif unsafe_kind == "merge":
        book["Data"].merge_cells("A2:B2")
        expected = "sheet_contains_merged_cells"
    else:
        book["Data"].protection.sheet = True
        expected = "protected_workbook"
    book.save(source)

    with pytest.raises(WorkbookStructureError) as caught:
        _apply(
            source,
            destination,
            action="insert_rows",
            sheet_key="sheet_0001",
            index=2,
        )
    assert caught.value.reason_code == expected
    assert not destination.exists()


# ---------------------------------------------------------------------------
# D-08: refuse lossy features and verify the whole workbook after save
# ---------------------------------------------------------------------------


def _add_feature(book, kind: str, tmp_path) -> None:
    sheet = book["Data"]
    if kind == "image":
        PIL = pytest.importorskip("PIL.Image")
        from openpyxl.drawing.image import Image as WorksheetImage

        image_path = tmp_path / "feature-dot.png"
        PIL.new("RGB", (4, 4), (255, 0, 0)).save(image_path)
        sheet.add_image(WorksheetImage(str(image_path)), "D4")
    elif kind == "validation":
        from openpyxl.worksheet.datavalidation import DataValidation

        validation = DataValidation(type="whole", operator="greaterThan", formula1="0")
        sheet.add_data_validation(validation)
        validation.add("A2")
    elif kind == "conditional_formatting":
        from openpyxl.formatting.rule import CellIsRule

        sheet.conditional_formatting.add(
            "A1:A3", CellIsRule(operator="greaterThan", formula=["0"])
        )
    elif kind == "comment":
        from openpyxl.comments import Comment

        sheet["A1"].comment = Comment("note", "tester")
    elif kind == "autofilter":
        sheet.auto_filter.ref = "A1:B3"
    elif kind == "hyperlink":
        sheet["A1"].hyperlink = "https://example.invalid/report"
    else:  # pragma: no cover - parametrization guard
        raise AssertionError(kind)


@pytest.mark.parametrize(
    ("feature", "expected"),
    [
        ("image", "sheet_contains_images"),
        ("validation", "sheet_contains_data_validations"),
        ("conditional_formatting", "sheet_contains_conditional_formatting"),
        ("comment", "sheet_contains_comments"),
        ("autofilter", "sheet_contains_autofilter"),
        ("hyperlink", "sheet_contains_hyperlinks"),
    ],
)
def test_row_operations_refuse_lossy_sheet_features(
    tmp_path, feature, expected
):
    from rag.spreadsheets.structure_editor import WorkbookStructureError

    source = tmp_path / "source.xlsx"
    destination = tmp_path / "edited.xlsx"
    _save_workbook(source)
    book = load_workbook(source)
    _add_feature(book, feature, tmp_path)
    book.save(source)
    book.close()

    with pytest.raises(WorkbookStructureError) as caught:
        _apply(
            source,
            destination,
            action="insert_rows",
            sheet_key="sheet_0001",
            index=2,
        )
    assert caught.value.reason_code == expected
    assert not destination.exists()


def test_add_sheet_is_verified_against_the_whole_workbook(tmp_path):
    source = tmp_path / "source.xlsx"
    destination = tmp_path / "edited.xlsx"
    _save_workbook(source)
    book = load_workbook(source)
    for feature in (
        "image",
        "validation",
        "conditional_formatting",
        "comment",
        "autofilter",
        "hyperlink",
    ):
        _add_feature(book, feature, tmp_path)
    book["Data"].merge_cells("A5:B5")
    book["Data"]["D5"] = "=1+1"
    book.save(source)
    book.close()

    _apply(source, destination, action="add_sheet", new_sheet_name="August")

    reopened = load_workbook(destination)
    data = reopened["Data"]
    assert reopened.sheetnames == ["Data", "Other", "August"]
    assert len(data._images) == 1
    assert data.data_validations.count == 1
    assert len(data.conditional_formatting) == 1
    assert data["A1"].comment is not None
    assert data["A1"].hyperlink is not None
    assert data.auto_filter.ref == "A1:B3"
    assert data["D5"].value == "=1+1"
    assert [str(value) for value in data.merged_cells.ranges] == ["A5:B5"]
    reopened.close()


def test_round_trip_validator_rejects_silent_feature_loss(tmp_path):
    from dataclasses import replace

    from rag.spreadsheets.structure_editor import (
        StructureChange,
        WorkbookStructureError,
        _snapshot_workbook,
        _validate_round_trip,
    )

    source = tmp_path / "source.xlsx"
    _save_workbook(source)
    book = load_workbook(source)
    before = _snapshot_workbook(book)
    book.create_sheet("August")
    saved = tmp_path / "saved.xlsx"
    book.save(saved)
    book.close()
    reopened = load_workbook(saved)
    after = _snapshot_workbook(reopened)
    reopened.close()
    change = StructureChange(
        action="add_sheet",
        sheet_key=None,
        sheet_name=None,
        index=None,
        count=1,
        new_sheet_name="August",
        affected_nonempty_cells=0,
    )

    _validate_round_trip(before, after, change)

    lost_comments = replace(
        after,
        sheets=tuple(
            replace(sheet, comments=sheet.comments + 1) if sheet.title == "Data" else sheet
            for sheet in after.sheets
        ),
    )
    with pytest.raises(WorkbookStructureError) as caught:
        _validate_round_trip(before, lost_comments, change)
    assert caught.value.reason_code == "workbook_validation_failed"

    renamed = replace(after, sheet_names=("Data", "August", "Other"))
    with pytest.raises(WorkbookStructureError):
        _validate_round_trip(before, renamed, change)


def test_shifted_sheet_validator_rejects_wrong_shift(tmp_path):
    from dataclasses import replace

    from rag.spreadsheets.structure_editor import (
        StructureChange,
        WorkbookStructureError,
        _snapshot_workbook,
        _validate_shifted_sheet,
    )

    source = tmp_path / "source.xlsx"
    _save_workbook(source)
    book = load_workbook(source)
    before = {sheet.title: sheet for sheet in _snapshot_workbook(book).sheets}
    book["Data"].insert_rows(2, 1)
    after = {sheet.title: sheet for sheet in _snapshot_workbook(book).sheets}
    book.close()
    change = StructureChange(
        action="insert_rows",
        sheet_key="sheet_0001",
        sheet_name="Data",
        index=2,
        count=1,
        new_sheet_name=None,
        affected_nonempty_cells=0,
    )

    _validate_shifted_sheet(before["Data"], after["Data"], change)

    unshifted = replace(after["Data"], values=before["Data"].values)
    with pytest.raises(WorkbookStructureError):
        _validate_shifted_sheet(before["Data"], unshifted, change)

    lost_style = replace(after["Data"], styles={})
    with pytest.raises(WorkbookStructureError):
        _validate_shifted_sheet(before["Data"], lost_style, change)
