"""Multi-sheet execution: the sheet dimension must survive planning.

Column keys are only unique *inside* one sheet (``rag.spreadsheets.parser``
numbers them per sheet), so a plan that references a bare ``c_000N`` cannot be
executed against a workbook whose sheets share keys.  These tests pin the fixed
behaviour: a sheet-scoped plan reads the named sheet, an unscoped colliding key
is an explicit error, and a question that names the same field in two sheets is
answered with a readable clarification instead of a guess.
"""

from __future__ import annotations

import pytest
from openpyxl import Workbook

from rag.spreadsheets.catalog import (
    CatalogColumn,
    CatalogDataset,
    CatalogSheet,
    SpreadsheetCatalog,
)
from rag.spreadsheets.coordinator import ComputedFact, SpreadsheetQueryCoordinator
from rag.spreadsheets.migration import dataset_schema_mirror
from rag.spreadsheets.parser import parse_spreadsheet
from rag.spreadsheets.planner import plan_spreadsheet_query
from rag.spreadsheets.query import (
    FilterClause,
    Metric,
    QueryScope,
    SortClause,
    StructuredQueryPlan,
    execute_multi_dataset_query,
    execute_structured_query,
)
from rag.spreadsheets.store import read_dataset_schema, resolve_dataset_path, write_dataset
from rag.user_db import UserDB

KB_ID = "kb_multi"
LENS_HEADER = ["镜头品牌", "型号", "遮光罩售价"]
LENS_ROWS = [["腾龙", "HA046", 320], ["适马", "LH582-01", 180], ["腾龙", "HA036", 260]]
BODY_HEADER = ["机身品牌", "型号", "电池单价", "备注"]
BODY_ROWS = [["佳能", "R5", 399, "-"], ["佳能", "R6", 299, "-"]]
# 型号 sits at c_0002 in *both* sheets; 遮光罩售价/电池单价 both sit at c_0003.
PRICE_QUESTION = "遮光罩售价最高是多少？"


def _write_two_sheet_workbook(path):
    book = Workbook()
    lens = book.active
    lens.title = "镜头"
    lens.append(LENS_HEADER)
    for row in LENS_ROWS:
        lens.append(row)
    body = book.create_sheet("机身")
    body.append(BODY_HEADER)
    for row in BODY_ROWS:
        body.append(row)
    book.save(path)
    return path


@pytest.fixture
def multi_sheet_dataset(tmp_path):
    """A real two-sheet dataset file built by the production writer."""
    source = _write_two_sheet_workbook(tmp_path / "multi.xlsx")
    path = resolve_dataset_path(tmp_path / "datasets", 1, 1)
    write_dataset(path, parse_spreadsheet(source, "xlsx", None))
    return path

def _scope(path):
    return QueryScope({1: path}, frozenset({1}))


def _catalog_from_file(path, dataset_id=1):
    """Catalog view of a dataset file, exactly as the users.db mirror exposes it."""
    sheets = tuple(
        CatalogSheet(
            sheet.sheet_key,
            sheet.name,
            sheet.row_count,
            tuple(
                CatalogColumn(
                    column.column_key,
                    column.display_name,
                    column.normalized_name,
                    column.inferred_type,
                )
                for column in sheet.columns
            ),
        )
        for sheet in read_dataset_schema(path)
    )
    return SpreadsheetCatalog(
        KB_ID, (CatalogDataset(dataset_id, 1, 1, "multi.xlsx", path, sheets),)
    )


def test_dataset_fixture_has_colliding_keys_across_sheets(multi_sheet_dataset):
    """Guard: the fixture really reproduces the production collision."""
    manifest = {
        sheet.sheet_key: {column.column_key: column.display_name for column in sheet.columns}
        for sheet in read_dataset_schema(multi_sheet_dataset)
    }
    assert set(manifest) == {"sheet_0001", "sheet_0002"}
    assert manifest["sheet_0001"]["c_0003"] == "遮光罩售价"
    assert manifest["sheet_0002"]["c_0003"] == "电池单价"
    assert manifest["sheet_0001"]["c_0002"] == manifest["sheet_0002"]["c_0002"] == "型号"


def test_sheet_scoped_plan_reads_only_that_sheet(multi_sheet_dataset):
    path = multi_sheet_dataset
    lens = StructuredQueryPlan(
        dataset_id=1,
        operation="aggregate",
        sheet_key="sheet_0001",
        metrics=(Metric(column_key="c_0003", function="max"),),
    )
    body = StructuredQueryPlan(
        dataset_id=1,
        operation="aggregate",
        sheet_key="sheet_0002",
        metrics=(Metric(column_key="c_0003", function="max"),),
    )

    lens_result = execute_structured_query(lens, _scope(path))
    body_result = execute_structured_query(body, _scope(path))

    assert lens_result.values[0].canonical == "320"
    assert body_result.values[0].canonical == "399"
    assert lens_result.rows_scanned == len(LENS_ROWS)
    assert body_result.rows_scanned == len(BODY_ROWS)
    assert lens_result.citations[0].sheet_key == "sheet_0001"
    assert body_result.citations[0].sheet_key == "sheet_0002"


def test_sheet_scoped_count_and_filter_use_the_named_table(multi_sheet_dataset):
    plan = StructuredQueryPlan(
        dataset_id=1,
        operation="aggregate",
        sheet_key="sheet_0001",
        metrics=(Metric(column_key="c_0002", function="count"),),
        filters=(FilterClause(column_key="c_0001", operator="eq", value="腾龙"),),
    )

    result = execute_structured_query(plan, _scope(multi_sheet_dataset))

    assert result.values[0].canonical == "2"
    assert result.rows_scanned == 2


def test_unqualified_colliding_key_is_an_explicit_error(multi_sheet_dataset):
    """c_0003 exists in both sheets: never guess one of them."""
    single = StructuredQueryPlan(
        dataset_id=1,
        operation="aggregate",
        metrics=(Metric(column_key="c_0003", function="count"),),
    )
    multi = StructuredQueryPlan(
        dataset_ids=(1,),
        operation="aggregate",
        metrics=(Metric(column_key="c_0003", function="count"),),
    )

    with pytest.raises(ValueError, match="unknown or ambiguous column key"):
        execute_structured_query(single, _scope(multi_sheet_dataset))
    with pytest.raises(ValueError, match="unknown or ambiguous column key"):
        execute_multi_dataset_query(multi, _scope(multi_sheet_dataset))


def test_unqualified_key_unique_to_one_sheet_still_resolves(multi_sheet_dataset):
    """Single-sheet behaviour is preserved: no unnecessary sheet scope required."""
    plan = StructuredQueryPlan(
        dataset_id=1,
        operation="aggregate",
        metrics=(Metric(column_key="c_0004", function="count"),),
    )

    result = execute_structured_query(plan, _scope(multi_sheet_dataset))

    assert result.values[0].canonical == str(len(BODY_ROWS))
    assert result.citations[0].sheet_key == "sheet_0002"


def test_lookup_scoped_to_sheet_returns_that_sheet_rows(multi_sheet_dataset):
    plan = StructuredQueryPlan(
        dataset_ids=(1,),
        operation="lookup",
        sheet_key="sheet_0002",
        select_columns=("c_0001", "c_0003"),
        order_by=(SortClause(column_key="c_0003", direction="desc"),),
    )

    result = execute_multi_dataset_query(plan, _scope(multi_sheet_dataset))

    assert [row.values for row in result.rows] == [
        {"c_0001": "佳能", "c_0003": 399},
        {"c_0001": "佳能", "c_0003": 299},
    ]
    assert {row.sheet_key for row in result.rows} == {"sheet_0002"}

    unscoped = StructuredQueryPlan(
        dataset_ids=(1,),
        operation="lookup",
        select_columns=("c_0001", "c_0003"),
    )
    with pytest.raises(ValueError, match="unknown or ambiguous column key"):
        execute_multi_dataset_query(unscoped, _scope(multi_sheet_dataset))


def test_planner_scopes_the_plan_to_the_matching_sheet(multi_sheet_dataset):
    planning = plan_spreadsheet_query(
        PRICE_QUESTION, _catalog_from_file(multi_sheet_dataset)
    )

    assert planning.kind == "planned"
    assert planning.plan is not None
    assert planning.plan.sheet_key == "sheet_0001"
    assert planning.plan.metrics[0].column_key == "c_0003"
    assert planning.plan.metrics[0].function == "max"

    result = execute_multi_dataset_query(planning.plan, _scope(multi_sheet_dataset))

    assert result.values[0].canonical == "320"
    assert result.citations[0].sheet_key == "sheet_0001"


def test_same_field_in_two_sheets_requires_a_sheet_clarification(multi_sheet_dataset):
    """型号 lives in both sheets (same key, same name): ask, do not pick one."""
    planning = plan_spreadsheet_query(
        "型号最高是多少？", _catalog_from_file(multi_sheet_dataset)
    )

    assert planning.kind == "clarification_required"
    assert planning.plan is None
    assert planning.clarification is not None
    assert planning.clarification.code == "ambiguous_sheet"
    assert set(planning.clarification.options) == {"镜头", "机身"}
    assert planning.clarification.question


def test_cross_sheet_grouping_is_rejected_with_a_readable_clarification(multi_sheet_dataset):
    """Grouping one sheet by a field of another sheet is not a supported join."""
    planning = plan_spreadsheet_query(
        "各机身品牌遮光罩售价合计是多少？", _catalog_from_file(multi_sheet_dataset)
    )

    assert planning.kind == "clarification_required"
    assert planning.plan is None
    assert planning.clarification is not None
    assert planning.clarification.code == "cross_sheet_unsupported"
    assert set(planning.clarification.options) == {"镜头", "机身"}
    assert planning.clarification.question


def test_grouping_inside_one_sheet_keeps_its_sheet_scope(multi_sheet_dataset):
    planning = plan_spreadsheet_query(
        "各镜头品牌遮光罩售价合计是多少？", _catalog_from_file(multi_sheet_dataset)
    )

    assert planning.kind == "planned"
    assert planning.plan is not None
    assert planning.plan.sheet_key == "sheet_0001"
    assert planning.plan.group_by == ("c_0001",)
    assert planning.plan.metrics[0].column_key == "c_0003"

    result = execute_multi_dataset_query(planning.plan, _scope(multi_sheet_dataset))

    assert [row.values["c_0003_sum"] for row in result.rows] == ["580", "180"]


def test_rejects_unknown_or_injected_sheet_keys(multi_sheet_dataset):
    injected = StructuredQueryPlan(
        dataset_id=1,
        operation="aggregate",
        sheet_key='sheet_0001"; DROP TABLE sheet_0002; --',
        metrics=(Metric(column_key="c_0003", function="count"),),
    )
    missing = StructuredQueryPlan(
        dataset_id=1,
        operation="aggregate",
        sheet_key="sheet_0009",
        metrics=(Metric(column_key="c_0003", function="count"),),
    )
    conflicting = StructuredQueryPlan(
        dataset_id=1,
        operation="aggregate",
        sheet_key="sheet_0001",
        metrics=(Metric(column_key="c_0003", function="count", sheet_key="sheet_0002"),),
    )
    cross_sheet_filter = StructuredQueryPlan(
        dataset_id=1,
        operation="aggregate",
        sheet_key="sheet_0001",
        metrics=(Metric(column_key="c_0003", function="count"),),
        filters=(
            FilterClause(
                column_key="c_0001", operator="eq", value="佳能", sheet_key="sheet_0002"
            ),
        ),
    )

    with pytest.raises(ValueError, match="invalid sheet or column key"):
        execute_structured_query(injected, _scope(multi_sheet_dataset))
    with pytest.raises(ValueError, match="unknown column key in sheet"):
        execute_structured_query(missing, _scope(multi_sheet_dataset))
    with pytest.raises(ValueError, match="conflicting sheet scope"):
        execute_structured_query(conflicting, _scope(multi_sheet_dataset))
    with pytest.raises(ValueError, match="cross-sheet filters are not supported"):
        execute_structured_query(cross_sheet_filter, _scope(multi_sheet_dataset))


def test_coordinator_executes_a_multi_sheet_dataset_end_to_end(tmp_path):
    """The real pipeline: dataset file -> mirror -> catalog -> planner -> execution."""
    db = UserDB(str(tmp_path / "users.sqlite3"))
    root = tmp_path / "datasets"
    owner_id = db.create_user("owner", "Password-123!")
    owner = db.get_user_by_id(owner_id)
    document_id = db.create_document_permission("multi.xlsx", KB_ID, owner_id)
    source = _write_two_sheet_workbook(tmp_path / "multi.xlsx")
    version_id = db.create_document_version(
        document_id, "multi.xlsx", str(source), "hash", owner_id
    )
    dataset_id = db.create_structured_dataset(document_id, version_id, "xlsx", "row_groups")
    path = resolve_dataset_path(root, document_id, version_id)
    write_dataset(path, parse_spreadsheet(source, "xlsx", None))
    db.replace_structured_dataset_schema(dataset_id, dataset_schema_mirror(path))
    db.activate_document_and_dataset(
        document_id=document_id,
        version_id=version_id,
        dataset_id=dataset_id,
        kb_id=KB_ID,
        filename="multi.xlsx",
    )

    outcome = SpreadsheetQueryCoordinator(db, root).query(
        PRICE_QUESTION, kb_id=KB_ID, user=owner
    )

    assert outcome.kind == "structured_result"
    assert outcome.rows_scanned == len(LENS_ROWS)
    assert outcome.computed_facts == (ComputedFact("c_0003 max", "320"),)
    assert outcome.citations[0].sheet_key == "sheet_0001"
    db.close()


def test_single_sheet_dataset_behaviour_is_unchanged(tmp_path):
    source = tmp_path / "single.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "销量"
    sheet.append(["品类", "销售额"])
    sheet.append(["相机", 30])
    sheet.append(["镜头", 20])
    book.save(source)
    path = resolve_dataset_path(tmp_path / "datasets", 2, 2)
    write_dataset(path, parse_spreadsheet(source, "xlsx", None))
    catalog = _catalog_from_file(path, dataset_id=2)

    planning = plan_spreadsheet_query("销售额合计是多少？", catalog)
    assert planning.kind == "planned"
    assert planning.plan is not None
    assert planning.plan.metrics[0].column_key == "c_0002"

    scoped = execute_multi_dataset_query(planning.plan, QueryScope({2: path}, frozenset({2})))
    unscoped = execute_multi_dataset_query(
        StructuredQueryPlan(
            dataset_id=2,
            operation="aggregate",
            metrics=(Metric(column_key="c_0002", function="sum"),),
        ),
        QueryScope({2: path}, frozenset({2})),
    )

    assert scoped.values[0].canonical == "50"
    assert scoped == unscoped
